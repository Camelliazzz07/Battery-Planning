# -*- coding: utf-8 -*-
"""问题2主脚本：滚动预测、风险备用、日计划优化、实际运行仿真并填充result2。

依赖：
    pip install numpy pandas scipy openpyxl

在 VS Code 中可直接运行当前文件。默认从仓库“附件”目录读取附件1.xlsx、
附件2.xlsx和“附件/附件5/result2.xlsx”，输出仍保存到问题二/q2_output：
    python q2_run_model.py

指定路径示例：
    python q2_run_model.py \
        --attachment1 附件1.xlsx \
        --attachment2 附件2.xlsx \
        --template result2.xlsx \
        --output-dir q2_output

建模口径：
1. 附件时间是10分钟区间的右端点；功率kW先乘1/6转为kWh。
2. 每天0:00只使用前一日及以前的实际数据预测当天，不使用未来实际值。
3. 计划购电按计划量收费；当天实际值只用于顺序运行仿真和紧急购电。
4. 主预测器是可解释的自适应历史集成；经验残差分位数用于风险计划和储能备用。
5. 优化使用MILP；储能按10分钟决策，result2中再汇总为4小时结果。
6. 内部始终使用自然日顺序，写入模板时才旋转为[第2点,...,第144点,第1点]。
"""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import sys
import warnings
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix


for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8", errors="replace")


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent
ATTACHMENT_DIR = PROJECT_ROOT / "附件"
TEMPLATE_DIR = ATTACHMENT_DIR / "附件5"

DT_HOURS = 1.0 / 6.0
T = 144
EXPECTED_MINUTES = np.arange(10, 1441, 10)
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
ETA_C = 0.90
ETA_D = 0.90
P_MAX_KW = 5000.0
STEP_MAX_KWH = P_MAX_KW * DT_HOURS
OUTPUT_START = pd.Timestamp("2025-02-01")
OUTPUT_END = pd.Timestamp("2025-12-31")
FOUR_HOUR_BLOCKS = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]


@dataclass(frozen=True)
class ModelConfig:
    plan_quantile: float = 0.80
    reserve_quantile: float = 0.95
    reserve_horizon_intervals: int = 24
    residual_lookback_days: int = 56
    min_residual_days: int = 7
    reserve_shortfall_penalty: float = 4.0
    terminal_soc_target_kwh: float = E_INITIAL
    terminal_soc_penalty: float = 0.65
    degradation_cost_per_kwh: float = 0.0
    total_spill_penalty: float = 1.0e-7
    mip_relative_gap: float = 1.0e-5
    emergency_tolerance_kwh: float = 1.0e-6
    use_down_reserve: bool = True


def parse_right_endpoint(value: object) -> int:
    """把附件表头转换为累计分钟；0:00+1转换为1440。"""
    if pd.isna(value):
        raise ValueError("时间值为空。")
    if isinstance(value, pd.Timestamp):
        return value.hour * 60 + value.minute
    if isinstance(value, datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float, np.integer, np.floating)):
        number = float(value)
        if 0 <= number <= 1:
            minute = int(round(number * 1440))
            return 1440 if minute == 0 else minute
        raise ValueError(f"无法识别数值型时间：{value!r}")
    text = str(value).strip().replace("：", ":")
    if text in {"0:00+1", "00:00+1", "24:00"}:
        return 1440
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?:\+0)?", text)
    if not match:
        raise ValueError(f"无法解析时间：{value!r}")
    hour, minute = int(match.group(1)), int(match.group(2))
    if hour > 23 or minute > 59:
        raise ValueError(f"时间超出范围：{value!r}")
    total = hour * 60 + minute
    return 1440 if total == 0 else total


def format_minute(minute: int) -> str:
    day_offset, minute_in_day = divmod(int(minute), 1440)
    hour, minute_part = divmod(minute_in_day, 60)
    suffix = f"+{day_offset}" if day_offset else ""
    return f"{hour}:{minute_part:02d}{suffix}"


def natural_interval_label(interval_index: int) -> str:
    start = interval_index * 10
    end = (interval_index + 1) * 10
    return f"{format_minute(start)}-{format_minute(end)}"


def _to_numeric_matrix(frame: pd.DataFrame, label: str) -> np.ndarray:
    numeric = frame.apply(pd.to_numeric, errors="coerce")
    missing = int(numeric.isna().sum().sum())
    if missing:
        raise ValueError(f"{label}存在{missing}个缺失或非数值单元格；主模型不自动插值。")
    values = numeric.to_numpy(dtype=float)
    if np.any(values < 0):
        raise ValueError(f"{label}存在负值。")
    return values


def read_attachment1(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    frame = pd.read_excel(path, engine="openpyxl")
    required = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    missing = [column for column in required if column not in frame.columns]
    if missing:
        raise KeyError(f"附件1缺少列：{missing}")
    frame = frame[required].copy()
    frame["右端点分钟"] = frame["时间"].map(parse_right_endpoint)
    frame = frame.sort_values("右端点分钟").reset_index(drop=True)
    if len(frame) != T or not np.array_equal(frame["右端点分钟"].to_numpy(), EXPECTED_MINUTES):
        raise ValueError("附件1必须包含按右端点排列的144个10分钟时段。")
    price = pd.to_numeric(frame["电价"], errors="raise").to_numpy(dtype=float)
    prior_load = pd.to_numeric(frame["小区负载"], errors="raise").to_numpy(dtype=float)
    prior_pv = pd.to_numeric(frame["光伏发电预测功率"], errors="raise").to_numpy(dtype=float)
    if np.any(price < 0) or np.any(prior_load < 0) or np.any(prior_pv < 0):
        raise ValueError("附件1存在负价格或负功率。")
    return price, prior_load, prior_pv


def read_attachment2(path: Path) -> tuple[pd.DatetimeIndex, np.ndarray, np.ndarray]:
    workbook = pd.ExcelFile(path, engine="openpyxl")
    required_sheets = ["小区负载", "光伏发电实际功率"]
    missing_sheets = [sheet for sheet in required_sheets if sheet not in workbook.sheet_names]
    if missing_sheets:
        raise KeyError(f"附件2缺少工作表：{missing_sheets}")
    load_wide = pd.read_excel(workbook, sheet_name="小区负载")
    pv_wide = pd.read_excel(workbook, sheet_name="光伏发电实际功率")
    if load_wide.shape != pv_wide.shape:
        raise ValueError("附件2负载表与光伏表形状不一致。")
    dates_load = pd.DatetimeIndex(
        pd.to_datetime(load_wide.iloc[:, 0], errors="raise").dt.normalize()
    )
    dates_pv = pd.DatetimeIndex(
        pd.to_datetime(pv_wide.iloc[:, 0], errors="raise").dt.normalize()
    )
    if not dates_load.equals(dates_pv):
        raise ValueError("附件2两个工作表的日期不一致。")
    if dates_load.has_duplicates:
        raise ValueError("附件2日期存在重复。")
    load_minutes = np.array([parse_right_endpoint(x) for x in load_wide.columns[1:]], dtype=int)
    pv_minutes = np.array([parse_right_endpoint(x) for x in pv_wide.columns[1:]], dtype=int)
    if not np.array_equal(load_minutes, pv_minutes):
        raise ValueError("附件2两个工作表的时间列不一致。")
    order = np.argsort(load_minutes)
    sorted_minutes = load_minutes[order]
    if len(sorted_minutes) != T or not np.array_equal(sorted_minutes, EXPECTED_MINUTES):
        raise ValueError("附件2每天必须包含完整的144个右端点。")
    load_kw = _to_numeric_matrix(load_wide.iloc[:, 1:].iloc[:, order], "附件2小区负载")
    pv_kw = _to_numeric_matrix(pv_wide.iloc[:, 1:].iloc[:, order], "附件2光伏实际功率")
    return dates_load, load_kw, pv_kw


def _candidate_for_day(
    values: np.ndarray,
    dates: pd.DatetimeIndex,
    day_index: int,
    name: str,
    prior_profile: np.ndarray,
) -> np.ndarray | None:
    if name == "附件1先验":
        return prior_profile.copy()
    if name == "前1天":
        return values[day_index - 1].copy() if day_index >= 1 else None
    if name == "前7天":
        return values[day_index - 7].copy() if day_index >= 7 else None
    if name == "近7天中位数":
        if day_index < 1:
            return None
        return np.median(values[max(0, day_index - 7):day_index], axis=0)
    if name == "同星期中位数":
        history_indices = [
            j for j in range(day_index)
            if dates[j].dayofweek == dates[day_index].dayofweek
        ][-8:]
        if not history_indices:
            return None
        return np.median(values[history_indices], axis=0)
    raise KeyError(name)


def adaptive_historical_forecast(
    values: np.ndarray,
    dates: pd.DatetimeIndex,
    day_index: int,
    prior_profile: np.ndarray,
    validation_days: int = 28,
) -> tuple[np.ndarray, dict[str, float]]:
    """只用过去数据，对多个历史基准按近期逆MAE加权。"""
    names = ["附件1先验", "前1天", "前7天", "近7天中位数", "同星期中位数"]
    current: dict[str, np.ndarray] = {}
    for name in names:
        candidate = _candidate_for_day(values, dates, day_index, name, prior_profile)
        if candidate is not None and np.isfinite(candidate).all():
            current[name] = candidate
    if not current:
        raise RuntimeError("没有可用的预测基准。")

    scores: dict[str, float] = {}
    validation_start = max(0, day_index - validation_days)
    for name in current:
        errors = []
        for historical_day in range(validation_start, day_index):
            candidate = _candidate_for_day(
                values, dates, historical_day, name, prior_profile
            )
            if candidate is not None:
                errors.append(float(np.mean(np.abs(values[historical_day] - candidate))))
        if errors:
            scores[name] = 1.0 / max(float(np.mean(errors)), 1.0e-6)
        else:
            scores[name] = 1.0e-6
    total_score = sum(scores.values())
    weights = {name: score / total_score for name, score in scores.items()}
    forecast = sum(weights[name] * current[name] for name in current)
    return np.asarray(forecast, dtype=float), weights


def _smooth_nonnegative(values: np.ndarray, window: int = 5) -> np.ndarray:
    return (
        pd.Series(np.maximum(values, 0.0))
        .rolling(window=window, center=True, min_periods=1)
        .median()
        .to_numpy(dtype=float)
    )


def build_risk_profiles(
    residual_history_kwh: list[np.ndarray],
    forecast_load_kwh: np.ndarray,
    forecast_pv_kwh: np.ndarray,
    config: ModelConfig,
) -> dict[str, np.ndarray]:
    """由历史整日净负荷残差构造计划分位数和额外尾部备用。"""
    if len(residual_history_kwh) < config.min_residual_days:
        base = 0.05 * forecast_load_kwh + 0.08 * forecast_pv_kwh
        plan_margin = _smooth_nonnegative(base)
        up_instant = _smooth_nonnegative(0.50 * base)
        down_instant = _smooth_nonnegative(0.50 * base) if config.use_down_reserve else np.zeros(T)
        scale = math.sqrt(config.reserve_horizon_intervals)
        up_cumulative = _smooth_nonnegative(up_instant * scale)
        down_cumulative = _smooth_nonnegative(down_instant * scale)
    else:
        errors = np.asarray(
            residual_history_kwh[-config.residual_lookback_days:], dtype=float
        )
        q_plan = np.quantile(errors, config.plan_quantile, axis=0)
        q_reserve = np.quantile(errors, config.reserve_quantile, axis=0)
        q_down = np.quantile(-errors, config.reserve_quantile, axis=0)
        plan_margin = _smooth_nonnegative(q_plan)
        up_instant = _smooth_nonnegative(q_reserve - np.maximum(q_plan, 0.0))
        down_instant = _smooth_nonnegative(q_down) if config.use_down_reserve else np.zeros(T)

        up_cumulative = np.zeros(T)
        down_cumulative = np.zeros(T)
        horizon = config.reserve_horizon_intervals
        for t in range(T):
            block = errors[:, t:min(T, t + horizon)]
            up_paths = np.maximum(np.cumsum(block, axis=1), 0.0).max(axis=1)
            q_plan_path = float(np.quantile(up_paths, config.plan_quantile))
            q_reserve_path = float(np.quantile(up_paths, config.reserve_quantile))
            up_cumulative[t] = max(0.0, q_reserve_path - q_plan_path)
            if config.use_down_reserve:
                down_paths = np.maximum(np.cumsum(-block, axis=1), 0.0).max(axis=1)
                down_cumulative[t] = max(
                    0.0, float(np.quantile(down_paths, config.reserve_quantile))
                )
        up_cumulative = _smooth_nonnegative(up_cumulative)
        down_cumulative = _smooth_nonnegative(down_cumulative)

    # 防止有限样本极端分位数使备用约束支配全部容量。
    up_instant = np.minimum(up_instant, 0.50 * STEP_MAX_KWH)
    down_instant = np.minimum(down_instant, 0.50 * STEP_MAX_KWH)
    up_energy_cap = 0.35 * (E_MAX - E_MIN) * ETA_D
    down_energy_cap = 0.35 * (E_MAX - E_MIN) / ETA_C
    up_cumulative = np.minimum(up_cumulative, up_energy_cap)
    down_cumulative = np.minimum(down_cumulative, down_energy_cap)
    return {
        "plan_margin": plan_margin,
        "up_instant": up_instant,
        "down_instant": down_instant,
        "up_cumulative": up_cumulative,
        "down_cumulative": down_cumulative,
    }


def solve_day_ahead_milp(
    price: np.ndarray,
    forecast_load_kwh: np.ndarray,
    forecast_pv_kwh: np.ndarray,
    start_soc_kwh: float,
    risk: dict[str, np.ndarray],
    config: ModelConfig,
) -> dict[str, np.ndarray | float | str]:
    """求解一天的风险感知MILP；计划量在0:00一次确定。"""
    if not (len(price) == len(forecast_load_kwh) == len(forecast_pv_kwh) == T):
        raise ValueError("日计划输入必须包含144个时段。")

    groups = ["q", "c", "d", "e", "w", "su", "sd", "z"]
    offsets = {name: i * T for i, name in enumerate(groups)}
    dev_pos = len(groups) * T
    dev_neg = dev_pos + 1
    n_vars = dev_neg + 1

    objective = np.zeros(n_vars)
    objective[offsets["q"]:offsets["q"] + T] = price
    objective[offsets["c"]:offsets["c"] + T] = config.degradation_cost_per_kwh
    objective[offsets["d"]:offsets["d"] + T] = config.degradation_cost_per_kwh
    objective[offsets["w"]:offsets["w"] + T] = config.total_spill_penalty
    objective[offsets["su"]:offsets["su"] + T] = config.reserve_shortfall_penalty
    objective[offsets["sd"]:offsets["sd"] + T] = config.reserve_shortfall_penalty
    objective[dev_pos] = config.terminal_soc_penalty
    objective[dev_neg] = config.terminal_soc_penalty

    lower = np.zeros(n_vars)
    upper = np.full(n_vars, np.inf)
    upper[offsets["c"]:offsets["c"] + T] = STEP_MAX_KWH
    upper[offsets["d"]:offsets["d"] + T] = STEP_MAX_KWH
    lower[offsets["e"]:offsets["e"] + T] = E_MIN
    upper[offsets["e"]:offsets["e"] + T] = E_MAX
    upper[offsets["z"]:offsets["z"] + T] = 1.0

    integrality = np.zeros(n_vars, dtype=int)
    integrality[offsets["z"]:offsets["z"] + T] = 1

    # 144条计划平衡 + 144条SOC递推 + 1条终端偏差定义。
    a_eq = lil_matrix((2 * T + 1, n_vars), dtype=float)
    b_eq = np.zeros(2 * T + 1)
    planned_load = forecast_load_kwh + risk["plan_margin"]
    for t in range(T):
        a_eq[t, offsets["q"] + t] = 1.0
        a_eq[t, offsets["c"] + t] = -1.0
        a_eq[t, offsets["d"] + t] = 1.0
        a_eq[t, offsets["w"] + t] = -1.0
        b_eq[t] = planned_load[t] - forecast_pv_kwh[t]

        row = T + t
        a_eq[row, offsets["e"] + t] = 1.0
        if t > 0:
            a_eq[row, offsets["e"] + t - 1] = -1.0
            b_eq[row] = 0.0
        else:
            b_eq[row] = start_soc_kwh
        a_eq[row, offsets["c"] + t] = -ETA_C
        a_eq[row, offsets["d"] + t] = 1.0 / ETA_D

    terminal_row = 2 * T
    a_eq[terminal_row, offsets["e"] + T - 1] = 1.0
    a_eq[terminal_row, dev_pos] = -1.0
    a_eq[terminal_row, dev_neg] = 1.0
    b_eq[terminal_row] = config.terminal_soc_target_kwh

    # 每时段：充电模式、放电模式、上备用SOC、下备用SOC，共4T条。
    a_ub = lil_matrix((4 * T, n_vars), dtype=float)
    b_ub = np.zeros(4 * T)
    charge_caps = np.maximum(0.0, STEP_MAX_KWH - risk["down_instant"])
    discharge_caps = np.maximum(0.0, STEP_MAX_KWH - risk["up_instant"])
    for t in range(T):
        # c <= charge_cap*z
        a_ub[t, offsets["c"] + t] = 1.0
        a_ub[t, offsets["z"] + t] = -charge_caps[t]
        b_ub[t] = 0.0
        # d <= discharge_cap*(1-z)
        row = T + t
        a_ub[row, offsets["d"] + t] = 1.0
        a_ub[row, offsets["z"] + t] = discharge_caps[t]
        b_ub[row] = discharge_caps[t]
        # E + su >= Emin + R_up/eta_d
        row = 2 * T + t
        a_ub[row, offsets["e"] + t] = -1.0
        a_ub[row, offsets["su"] + t] = -1.0
        b_ub[row] = -(E_MIN + risk["up_cumulative"][t] / ETA_D)
        # E - sd <= Emax - eta_c*R_down
        row = 3 * T + t
        a_ub[row, offsets["e"] + t] = 1.0
        a_ub[row, offsets["sd"] + t] = -1.0
        b_ub[row] = E_MAX - ETA_C * risk["down_cumulative"][t]

    constraints = [
        LinearConstraint(a_eq.tocsr(), b_eq, b_eq),
        LinearConstraint(a_ub.tocsr(), -np.inf, b_ub),
    ]
    result = milp(
        c=objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=constraints,
        options={"mip_rel_gap": config.mip_relative_gap, "presolve": True},
    )
    if not result.success or result.x is None:
        raise RuntimeError(f"日计划MILP求解失败：status={result.status}, message={result.message}")

    values = result.x
    output: dict[str, np.ndarray | float | str] = {
        name: values[offsets[name]:offsets[name] + T].copy() for name in groups
    }
    output["objective"] = float(result.fun)
    output["solver_message"] = str(result.message)
    output["terminal_deviation"] = float(values[dev_pos] + values[dev_neg])
    return output


def simulate_actual_operation(
    plan_q: np.ndarray,
    plan_c: np.ndarray,
    plan_d: np.ndarray,
    actual_load_kwh: np.ndarray,
    actual_pv_kwh: np.ndarray,
    start_soc_kwh: float,
) -> dict[str, np.ndarray]:
    """计划跟踪型因果控制：先跟随计划，再用可用储能处理已发生的误差。"""
    actual_c = np.zeros(T)
    actual_d = np.zeros(T)
    actual_e = np.zeros(T)
    emergency = np.zeros(T)
    spill = np.zeros(T)
    pv_curtail = np.zeros(T)
    grid_unused = np.zeros(T)
    start_soc = np.zeros(T)
    e = float(start_soc_kwh)

    for t in range(T):
        start_soc[t] = e
        max_charge = min(STEP_MAX_KWH, max(0.0, (E_MAX - e) / ETA_C))
        max_discharge = min(STEP_MAX_KWH, max(0.0, ETA_D * (e - E_MIN)))
        c = min(max(float(plan_c[t]), 0.0), max_charge)
        d = min(max(float(plan_d[t]), 0.0), max_discharge)
        if c > 1.0e-8 and d > 1.0e-8:
            # 求解容差下只保留净作用较大的方向。
            if c >= d:
                d = 0.0
            else:
                c = 0.0

        balance = plan_q[t] + actual_pv_kwh[t] + d - actual_load_kwh[t] - c
        if balance < 0:
            shortage = -balance
            if c > 0:
                reduction = min(c, shortage)
                c -= reduction
                shortage -= reduction
            if shortage > 0 and c <= 1.0e-8:
                max_discharge = min(STEP_MAX_KWH, max(0.0, ETA_D * (e - E_MIN)))
                addition = min(max_discharge - d, shortage)
                if addition > 0:
                    d += addition
                    shortage -= addition
            emergency[t] = max(0.0, shortage)
        else:
            surplus = balance
            if d > 0:
                reduction = min(d, surplus)
                d -= reduction
                surplus -= reduction
            if surplus > 0 and d <= 1.0e-8:
                max_charge = min(STEP_MAX_KWH, max(0.0, (E_MAX - e) / ETA_C))
                addition = min(max_charge - c, surplus)
                if addition > 0:
                    c += addition
                    surplus -= addition
            spill[t] = max(0.0, surplus)

        e = e + ETA_C * c - d / ETA_D
        if e < E_MIN - 1.0e-5 or e > E_MAX + 1.0e-5:
            raise RuntimeError(f"实际运行SOC越界：t={t + 1}, SOC={e}")
        e = float(np.clip(e, E_MIN, E_MAX))
        actual_c[t], actual_d[t], actual_e[t] = c, d, e
        pv_curtail[t] = min(actual_pv_kwh[t], spill[t])
        grid_unused[t] = max(0.0, spill[t] - actual_pv_kwh[t])

    return {
        "start_soc": start_soc,
        "charge": actual_c,
        "discharge": actual_d,
        "soc": actual_e,
        "emergency": emergency,
        "spill": spill,
        "pv_curtail": pv_curtail,
        "grid_unused": grid_unused,
    }


def aggregate_emergency_events(
    dates: pd.DatetimeIndex,
    emergency_matrix: np.ndarray,
    tolerance: float,
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for day_index, date in enumerate(dates):
        values = emergency_matrix[day_index]
        positive = values > tolerance
        t = 0
        found = False
        while t < T:
            if not positive[t]:
                t += 1
                continue
            start = t
            while t + 1 < T and positive[t + 1]:
                t += 1
            end = t
            rows.append(
                {
                    "日期": date,
                    "购电时间段": f"{format_minute(start * 10)}-{format_minute((end + 1) * 10)}",
                    "购电量": float(values[start:end + 1].sum()),
                }
            )
            found = True
            t += 1
        if not found:
            rows.append({"日期": date, "购电时间段": None, "购电量": None})
    return pd.DataFrame(rows)


def copy_row_style(source_ws, target_ws, source_row: int, target_row: int) -> None:
    for column in range(1, source_ws.max_column + 1):
        source = source_ws.cell(source_row, column)
        target = target_ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _clear_existing_data(ws) -> None:
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def write_result2(
    template_path: Path,
    output_path: Path,
    output_dates: pd.DatetimeIndex,
    price: np.ndarray,
    plan_q_matrix: np.ndarray,
    actual_charge_matrix: np.ndarray,
    actual_discharge_matrix: np.ndarray,
    start_soc_matrix: np.ndarray,
    actual_soc_matrix: np.ndarray,
    emergency_events: pd.DataFrame,
) -> None:
    """保留模板工作表名称和主要样式，替换示例/省略号并填入完整结果。"""
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path)
    required = ["计划购电量", "充放电量", "紧急购电量"]
    missing = [sheet for sheet in required if sheet not in workbook.sheetnames]
    if missing:
        raise KeyError(f"result2模板缺少工作表：{missing}")

    # 计划购电量工作表
    ws_plan = workbook["计划购电量"]
    headers = [ws_plan.cell(1, col).value for col in range(1, ws_plan.max_column + 1)]
    try:
        total_col = headers.index("全天购电量") + 1
        cost_col = headers.index("全天购电费") + 1
    except ValueError as exc:
        raise ValueError("result2计划购电量工作表缺少合计列。") from exc
    interval_cols = list(range(2, total_col))
    if len(interval_cols) != T:
        raise ValueError(f"result2计划购电量工作表应有144个时段列，实际{len(interval_cols)}个。")
    plan_style_source = 2 if ws_plan.max_row >= 2 else 1
    plan_styles = [
        copy(ws_plan.cell(plan_style_source, c)._style)
        for c in range(1, ws_plan.max_column + 1)
    ]
    _clear_existing_data(ws_plan)
    for day_index, date in enumerate(output_dates):
        row = day_index + 2
        for column in range(1, ws_plan.max_column + 1):
            ws_plan.cell(row, column)._style = copy(plan_styles[column - 1])
        ws_plan.cell(row, 1, date.to_pydatetime())
        ws_plan.cell(row, 1).number_format = "yyyy-mm-dd"
        natural = plan_q_matrix[day_index]
        rotated = np.r_[natural[1:], natural[:1]]
        for col, value in zip(interval_cols, rotated):
            ws_plan.cell(row, col, float(max(0.0, value)))
            ws_plan.cell(row, col).number_format = "0.000"
        ws_plan.cell(row, total_col, float(natural.sum()))
        ws_plan.cell(row, cost_col, float(np.dot(price, natural)))
        ws_plan.cell(row, total_col).number_format = "0.000"
        ws_plan.cell(row, cost_col).number_format = "0.00"
    ws_plan.freeze_panes = "B2"

    # 充放电量工作表：每个日期固定6行。
    ws_storage = workbook["充放电量"]
    storage_style_source = 2 if ws_storage.max_row >= 2 else 1
    source_styles = [copy(ws_storage.cell(storage_style_source, c)._style) for c in range(1, 7)]
    _clear_existing_data(ws_storage)
    target_row = 2
    for day_index, date in enumerate(output_dates):
        for block_index, block_label in enumerate(FOUR_HOUR_BLOCKS):
            for column in range(1, 7):
                ws_storage.cell(target_row, column)._style = copy(source_styles[column - 1])
            start, end = block_index * 24, (block_index + 1) * 24
            ws_storage.cell(target_row, 1, date.to_pydatetime() if block_index == 0 else None)
            if block_index == 0:
                ws_storage.cell(target_row, 1).number_format = "yyyy-mm-dd"
            ws_storage.cell(target_row, 2, block_label)
            ws_storage.cell(target_row, 3, float(actual_charge_matrix[day_index, start:end].sum()))
            ws_storage.cell(target_row, 4, float(actual_discharge_matrix[day_index, start:end].sum()))
            if block_index == 0:
                ws_storage.cell(target_row, 5, "0:00")
                ws_storage.cell(target_row, 6, float(start_soc_matrix[day_index, 0]))
            elif block_index == 1:
                ws_storage.cell(target_row, 5, "24:00")
                ws_storage.cell(target_row, 6, float(actual_soc_matrix[day_index, -1]))
            for column in (3, 4, 6):
                ws_storage.cell(target_row, column).number_format = "0.000"
            target_row += 1
    ws_storage.freeze_panes = "A2"

    # 紧急购电量工作表：有事件则逐事件列出，无事件仍保留该日期一行。
    ws_emergency = workbook["紧急购电量"]
    emergency_style_source = 2 if ws_emergency.max_row >= 2 else 1
    emergency_styles = [copy(ws_emergency.cell(emergency_style_source, c)._style) for c in range(1, 4)]
    _clear_existing_data(ws_emergency)
    target_row = 2
    for date, group in emergency_events.groupby("日期", sort=True):
        for item_index, (_, item) in enumerate(group.iterrows()):
            for column in range(1, 4):
                ws_emergency.cell(target_row, column)._style = copy(emergency_styles[column - 1])
            ws_emergency.cell(target_row, 1, pd.Timestamp(date).to_pydatetime() if item_index == 0 else None)
            if item_index == 0:
                ws_emergency.cell(target_row, 1).number_format = "yyyy-mm-dd"
            ws_emergency.cell(target_row, 2, item["购电时间段"])
            amount = item["购电量"]
            ws_emergency.cell(target_row, 3, None if pd.isna(amount) else float(amount))
            ws_emergency.cell(target_row, 3).number_format = "0.000"
            target_row += 1
    ws_emergency.freeze_panes = "A2"

    # 删除可能残留的省略号颜色，并确保工作表仍按模板顺序排列。
    for ws in (ws_storage, ws_emergency):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value in {"……", "…", "⁝"}:
                    cell.value = None
                    cell.fill = PatternFill(fill_type=None)

    workbook.save(output_path)


def run_model(args: argparse.Namespace) -> dict[str, Path]:
    config = ModelConfig(
        plan_quantile=args.plan_quantile,
        reserve_quantile=args.reserve_quantile,
        reserve_horizon_intervals=args.reserve_horizon,
        residual_lookback_days=args.residual_lookback,
        reserve_shortfall_penalty=args.reserve_shortfall_penalty,
        terminal_soc_target_kwh=args.terminal_soc_target,
        terminal_soc_penalty=args.terminal_soc_penalty,
        degradation_cost_per_kwh=args.degradation_cost,
        use_down_reserve=not args.disable_down_reserve,
    )
    if not (0.50 <= config.plan_quantile < config.reserve_quantile < 1.0):
        raise ValueError("必须满足0.50 <= plan_quantile < reserve_quantile < 1。")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    price, prior_load_kw, prior_pv_kw = read_attachment1(args.attachment1)
    dates, load_kw, pv_kw = read_attachment2(args.attachment2)
    if dates.min() != pd.Timestamp("2025-01-01") or dates.max() != OUTPUT_END:
        warnings.warn("附件2日期范围不是完整的2025年；请核对输出日期。", RuntimeWarning)

    load_kwh = load_kw * DT_HOURS
    pv_kwh = pv_kw * DT_HOURS
    n_days = len(dates)
    shape = (n_days, T)
    arrays = {
        name: np.zeros(shape, dtype=float)
        for name in [
            "forecast_load_kw", "forecast_pv_kw", "forecast_load_kwh", "forecast_pv_kwh",
            "plan_margin", "up_instant", "down_instant", "up_cumulative", "down_cumulative",
            "plan_q", "plan_c", "plan_d", "plan_soc", "plan_spill",
            "reserve_shortfall_up", "reserve_shortfall_down",
            "actual_start_soc", "actual_c", "actual_d", "actual_soc", "emergency",
            "spill", "pv_curtail", "grid_unused",
        ]
    }
    objectives = np.zeros(n_days)
    terminal_deviation = np.zeros(n_days)
    load_weight_records: list[dict[str, object]] = []
    pv_weight_records: list[dict[str, object]] = []
    residual_history: list[np.ndarray] = []
    current_soc = E_INITIAL

    print("[开始] 逐日滚动预测、优化和实际运行仿真")
    for day_index, date in enumerate(dates):
        forecast_load_kw, load_weights = adaptive_historical_forecast(
            load_kw, dates, day_index, prior_load_kw
        )
        forecast_pv_kw, pv_weights = adaptive_historical_forecast(
            pv_kw, dates, day_index, prior_pv_kw
        )
        forecast_load_kw = np.maximum(forecast_load_kw, 0.0)
        forecast_pv_kw = np.maximum(forecast_pv_kw, 0.0)
        forecast_load_kwh = forecast_load_kw * DT_HOURS
        forecast_pv_kwh = forecast_pv_kw * DT_HOURS
        risk = build_risk_profiles(
            residual_history, forecast_load_kwh, forecast_pv_kwh, config
        )
        plan = solve_day_ahead_milp(
            price, forecast_load_kwh, forecast_pv_kwh, current_soc, risk, config
        )
        actual = simulate_actual_operation(
            np.asarray(plan["q"]),
            np.asarray(plan["c"]),
            np.asarray(plan["d"]),
            load_kwh[day_index],
            pv_kwh[day_index],
            current_soc,
        )

        arrays["forecast_load_kw"][day_index] = forecast_load_kw
        arrays["forecast_pv_kw"][day_index] = forecast_pv_kw
        arrays["forecast_load_kwh"][day_index] = forecast_load_kwh
        arrays["forecast_pv_kwh"][day_index] = forecast_pv_kwh
        for name in ["plan_margin", "up_instant", "down_instant", "up_cumulative", "down_cumulative"]:
            arrays[name][day_index] = risk[name]
        arrays["plan_q"][day_index] = np.asarray(plan["q"])
        arrays["plan_c"][day_index] = np.asarray(plan["c"])
        arrays["plan_d"][day_index] = np.asarray(plan["d"])
        arrays["plan_soc"][day_index] = np.asarray(plan["e"])
        arrays["plan_spill"][day_index] = np.asarray(plan["w"])
        arrays["reserve_shortfall_up"][day_index] = np.asarray(plan["su"])
        arrays["reserve_shortfall_down"][day_index] = np.asarray(plan["sd"])
        arrays["actual_start_soc"][day_index] = actual["start_soc"]
        arrays["actual_c"][day_index] = actual["charge"]
        arrays["actual_d"][day_index] = actual["discharge"]
        arrays["actual_soc"][day_index] = actual["soc"]
        arrays["emergency"][day_index] = actual["emergency"]
        arrays["spill"][day_index] = actual["spill"]
        arrays["pv_curtail"][day_index] = actual["pv_curtail"]
        arrays["grid_unused"][day_index] = actual["grid_unused"]
        objectives[day_index] = float(plan["objective"])
        terminal_deviation[day_index] = float(plan["terminal_deviation"])
        current_soc = float(actual["soc"][-1])

        net_residual = (load_kwh[day_index] - forecast_load_kwh) - (
            pv_kwh[day_index] - forecast_pv_kwh
        )
        residual_history.append(net_residual)
        for name, weight in load_weights.items():
            load_weight_records.append({"日期": date, "变量": "负载", "基准": name, "权重": weight})
        for name, weight in pv_weights.items():
            pv_weight_records.append({"日期": date, "变量": "光伏", "基准": name, "权重": weight})

        if day_index == 0 or (day_index + 1) % 30 == 0 or day_index == n_days - 1:
            print(
                f"  {date.date()} 进度 {day_index + 1}/{n_days}，"
                f"SOC={current_soc:.1f} kWh，紧急购电={actual['emergency'].sum():.1f} kWh"
            )

    output_mask = (dates >= OUTPUT_START) & (dates <= OUTPUT_END)
    output_dates = dates[output_mask]
    output_indices = np.flatnonzero(output_mask)

    # 10分钟长表，供独立诊断脚本使用。
    records: list[pd.DataFrame] = []
    interval_ids = np.arange(1, T + 1)
    interval_labels = [natural_interval_label(i) for i in range(T)]
    for day_index in output_indices:
        net_actual_kwh = load_kwh[day_index] - pv_kwh[day_index]
        net_forecast_kwh = arrays["forecast_load_kwh"][day_index] - arrays["forecast_pv_kwh"][day_index]
        frame = pd.DataFrame(
            {
                "日期": dates[day_index],
                "区间编号": interval_ids,
                "区间": interval_labels,
                "右端点分钟": EXPECTED_MINUTES,
                "电价_元每kWh": price,
                "实际负载_kW": load_kw[day_index],
                "实际光伏_kW": pv_kw[day_index],
                "预测负载_kW": arrays["forecast_load_kw"][day_index],
                "预测光伏_kW": arrays["forecast_pv_kw"][day_index],
                "实际负载_kWh": load_kwh[day_index],
                "实际光伏_kWh": pv_kwh[day_index],
                "预测净负荷_kWh": net_forecast_kwh,
                "实际净负荷_kWh": net_actual_kwh,
                "净负荷预测误差_kWh": net_actual_kwh - net_forecast_kwh,
                "计划风险增量_kWh": arrays["plan_margin"][day_index],
                "净负荷预测下界_kWh": net_forecast_kwh - arrays["down_instant"][day_index],
                "净负荷预测上界_kWh": net_forecast_kwh + arrays["plan_margin"][day_index] + arrays["up_instant"][day_index],
                "计划购电量_kWh": arrays["plan_q"][day_index],
                "计划充电量_kWh": arrays["plan_c"][day_index],
                "计划放电量_kWh": arrays["plan_d"][day_index],
                "计划SOC_kWh": arrays["plan_soc"][day_index],
                "计划剩余电量_kWh": arrays["plan_spill"][day_index],
                "实际期初SOC_kWh": arrays["actual_start_soc"][day_index],
                "实际充电量_kWh": arrays["actual_c"][day_index],
                "实际放电量_kWh": arrays["actual_d"][day_index],
                "实际SOC_kWh": arrays["actual_soc"][day_index],
                "紧急购电量_kWh": arrays["emergency"][day_index],
                "总剩余电量_kWh": arrays["spill"][day_index],
                "弃光量_kWh": arrays["pv_curtail"][day_index],
                "未利用外网电量_kWh": arrays["grid_unused"][day_index],
                "正向瞬时备用_kWh": arrays["up_instant"][day_index],
                "反向瞬时备用_kWh": arrays["down_instant"][day_index],
                "正向累计备用_kWh": arrays["up_cumulative"][day_index],
                "反向累计备用_kWh": arrays["down_cumulative"][day_index],
                "正向备用缺口_kWh": arrays["reserve_shortfall_up"][day_index],
                "反向备用缺口_kWh": arrays["reserve_shortfall_down"][day_index],
            }
        )
        records.append(frame)
    timeseries = pd.concat(records, ignore_index=True)
    timeseries_path = args.output_dir / "q2_timeseries.csv"
    timeseries.to_csv(timeseries_path, index=False, encoding="utf-8-sig")

    daily_rows = []
    for day_index in output_indices:
        plan_cost = float(np.dot(price, arrays["plan_q"][day_index]))
        emergency_cost = float(np.dot(5.0 * price, arrays["emergency"][day_index]))
        available_pv = float(pv_kwh[day_index].sum())
        curtailed = float(arrays["pv_curtail"][day_index].sum())
        daily_rows.append(
            {
                "日期": dates[day_index],
                "预测负载MAE_kW": float(np.mean(np.abs(load_kw[day_index] - arrays["forecast_load_kw"][day_index]))),
                "预测光伏MAE_kW": float(np.mean(np.abs(pv_kw[day_index] - arrays["forecast_pv_kw"][day_index]))),
                "预测净负荷MAE_kWh": float(np.mean(np.abs((load_kwh[day_index] - pv_kwh[day_index]) - (arrays["forecast_load_kwh"][day_index] - arrays["forecast_pv_kwh"][day_index])))),
                "计划购电量_kWh": float(arrays["plan_q"][day_index].sum()),
                "计划购电费_元": plan_cost,
                "紧急购电量_kWh": float(arrays["emergency"][day_index].sum()),
                "紧急购电费_元": emergency_cost,
                "总购电费_元": plan_cost + emergency_cost,
                "弃光量_kWh": curtailed,
                "弃光率": curtailed / available_pv if available_pv > 0 else 0.0,
                "未利用外网电量_kWh": float(arrays["grid_unused"][day_index].sum()),
                "期初SOC_kWh": float(arrays["actual_start_soc"][day_index, 0]),
                "期末SOC_kWh": float(arrays["actual_soc"][day_index, -1]),
                "正向备用缺口_kWh": float(arrays["reserve_shortfall_up"][day_index].sum()),
                "反向备用缺口_kWh": float(arrays["reserve_shortfall_down"][day_index].sum()),
                "优化目标值": float(objectives[day_index]),
                "计划终端SOC偏差_kWh": float(terminal_deviation[day_index]),
            }
        )
    daily = pd.DataFrame(daily_rows)
    daily_path = args.output_dir / "q2_daily_summary.csv"
    daily.to_csv(daily_path, index=False, encoding="utf-8-sig")

    weights = pd.DataFrame(load_weight_records + pv_weight_records)
    weights = weights[weights["日期"].between(OUTPUT_START, OUTPUT_END)]
    weights_path = args.output_dir / "q2_forecast_weights.csv"
    weights.to_csv(weights_path, index=False, encoding="utf-8-sig")

    emergency_events = aggregate_emergency_events(
        output_dates,
        arrays["emergency"][output_mask],
        config.emergency_tolerance_kwh,
    )
    emergency_path = args.output_dir / "q2_emergency_events.csv"
    emergency_events.to_csv(emergency_path, index=False, encoding="utf-8-sig")

    result_path = args.output_dir / "result2_filled.xlsx"
    write_result2(
        args.template,
        result_path,
        output_dates,
        price,
        arrays["plan_q"][output_mask],
        arrays["actual_c"][output_mask],
        arrays["actual_d"][output_mask],
        arrays["actual_start_soc"][output_mask],
        arrays["actual_soc"][output_mask],
        emergency_events,
    )

    metadata = {
        "model": "causal adaptive ensemble + empirical quantile reserve + day-ahead MILP + causal plan-following operation",
        "config": asdict(config),
        "data_range": [str(dates.min().date()), str(dates.max().date())],
        "output_range": [str(output_dates.min().date()), str(output_dates.max().date())],
        "right_endpoint_rule": "0:10 means 0:00-0:10",
        "template_rotation": "[point 2, ..., point 144, point 1]",
        "units": {"source_power": "kW", "decision_energy": "kWh", "dt_hours": DT_HOURS},
        "cost_rule": "planned purchase is billed by planned quantity; emergency price is 5x",
        "cold_start": "Attachment 1 load/PV profiles are used as prior; January is causal warm-up",
        "output_files": [
            result_path.name, timeseries_path.name, daily_path.name,
            weights_path.name, emergency_path.name,
        ],
    }
    metadata_path = args.output_dir / "q2_run_metadata.json"
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n[完成]")
    print(f"  result2：{result_path.resolve()}")
    print(f"  10分钟明细：{timeseries_path.resolve()}")
    print(f"  每日汇总：{daily_path.resolve()}")
    print(f"  计划购电费：{daily['计划购电费_元'].sum():,.2f} 元")
    print(f"  紧急购电费：{daily['紧急购电费_元'].sum():,.2f} 元")
    print(f"  紧急购电量：{daily['紧急购电量_kWh'].sum():,.2f} kWh")
    print(f"  全期弃光率：{daily['弃光量_kWh'].sum() / (pv_kwh[output_mask].sum()):.2%}")
    return {
        "result2": result_path,
        "timeseries": timeseries_path,
        "daily": daily_path,
        "weights": weights_path,
        "emergency": emergency_path,
        "metadata": metadata_path,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="问题2滚动预测、优化、实际仿真与result2填充")
    parser.add_argument(
        "--attachment1", type=Path, default=ATTACHMENT_DIR / "附件1.xlsx"
    )
    parser.add_argument(
        "--attachment2", type=Path, default=ATTACHMENT_DIR / "附件2.xlsx"
    )
    parser.add_argument(
        "--template", type=Path, default=TEMPLATE_DIR / "result2.xlsx"
    )
    parser.add_argument(
        "--output-dir", type=Path, default=SCRIPT_DIR / "q2_output"
    )
    parser.add_argument("--plan-quantile", type=float, default=0.80)
    parser.add_argument("--reserve-quantile", type=float, default=0.95)
    parser.add_argument("--reserve-horizon", type=int, default=24, help="累计备用观察时段数，24表示4小时")
    parser.add_argument("--residual-lookback", type=int, default=56)
    parser.add_argument("--reserve-shortfall-penalty", type=float, default=4.0)
    parser.add_argument("--terminal-soc-target", type=float, default=E_INITIAL)
    parser.add_argument("--terminal-soc-penalty", type=float, default=0.65)
    parser.add_argument("--degradation-cost", type=float, default=0.0, help="有文献依据时再填写元/kWh吞吐成本")
    parser.add_argument("--disable-down-reserve", action="store_true", help="关闭用于吸收额外光伏的反向备用")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    for path, label in [
        (args.attachment1, "附件1"),
        (args.attachment2, "附件2"),
        (args.template, "result2模板"),
    ]:
        if not path.exists():
            raise FileNotFoundError(f"找不到{label}：{path.resolve()}")
    run_model(args)


if __name__ == "__main__":
    main()
