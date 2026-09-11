# -*- coding: utf-8 -*-
"""问题2独立诊断脚本：预测、风险、物理约束、费用、result2和可视化检验。

先运行 q2_run_model.py，再在 VS Code 中直接运行当前文件即可。默认读取
本脚本所在目录的q2_output，并将诊断结果保存到q2_diagnostics。

依赖：
    pip install numpy pandas matplotlib seaborn openpyxl

本脚本只读取主脚本的结果，不参与计划制定，不会把诊断结果反馈给模型。
"""

from __future__ import annotations

import argparse
import os
import sys
import warnings
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl import load_workbook


for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8", errors="replace")


SCRIPT_DIR = Path(__file__).resolve().parent

T = 144
ETA_C = 0.90
ETA_D = 0.90
E_MIN = 1200.0
E_MAX = 10800.0
STEP_MAX_KWH = 5000.0 / 6.0
EXPECTED_DAYS = 334
KEY_DATES = pd.to_datetime(["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"])


def configure_plot_style() -> None:
    sns.set_theme(style="whitegrid", context="notebook")
    custom_font = os.environ.get("EDA_FONT_PATH")
    custom_name = None
    if custom_font and Path(custom_font).exists():
        font_manager.fontManager.addfont(custom_font)
        custom_name = font_manager.FontProperties(fname=custom_font).get_name()
    installed = {item.name for item in font_manager.fontManager.ttflist}
    candidates = [
        "Microsoft YaHei", "SimHei", "Noto Sans SC", "Noto Sans CJK SC",
        "Source Han Sans SC", "Arial Unicode MS", "DejaVu Sans",
    ]
    selected = custom_name or next((x for x in candidates if x in installed), "DejaVu Sans")
    plt.rcParams["font.sans-serif"] = [selected] + candidates
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 120
    if selected == "DejaVu Sans":
        print("[字体提示] 未检测到中文字体；可安装SimHei/微软雅黑或设置EDA_FONT_PATH。")
        warnings.filterwarnings("ignore", message=r"Glyph .* missing from font")


def load_outputs(input_dir: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    paths = {
        "timeseries": input_dir / "q2_timeseries.csv",
        "daily": input_dir / "q2_daily_summary.csv",
        "events": input_dir / "q2_emergency_events.csv",
    }
    missing = [str(path) for path in paths.values() if not path.exists()]
    if missing:
        raise FileNotFoundError(f"缺少主脚本输出：{missing}")
    ts = pd.read_csv(paths["timeseries"])
    daily = pd.read_csv(paths["daily"])
    events = pd.read_csv(paths["events"])
    ts["日期"] = pd.to_datetime(ts["日期"])
    daily["日期"] = pd.to_datetime(daily["日期"])
    events["日期"] = pd.to_datetime(events["日期"])
    return ts, daily, events


def scalar_metrics(actual: np.ndarray, forecast: np.ndarray, prefix: str) -> dict[str, object]:
    error = forecast - actual
    return {
        "变量": prefix,
        "样本数": int(len(error)),
        "MAE": float(np.mean(np.abs(error))),
        "RMSE": float(np.sqrt(np.mean(error**2))),
        "Bias_预测减实际": float(np.mean(error)),
        "P95绝对误差": float(np.quantile(np.abs(error), 0.95)),
    }


def forecast_diagnostics(ts: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    metrics = [
        scalar_metrics(ts["实际负载_kW"].to_numpy(), ts["预测负载_kW"].to_numpy(), "负载功率_kW"),
        scalar_metrics(ts["实际光伏_kW"].to_numpy(), ts["预测光伏_kW"].to_numpy(), "光伏功率_kW"),
        scalar_metrics(ts["实际净负荷_kWh"].to_numpy(), ts["预测净负荷_kWh"].to_numpy(), "净负荷电量_kWh"),
    ]
    metric_frame = pd.DataFrame(metrics)
    ts = ts.copy()
    ts["区间覆盖"] = ts["实际净负荷_kWh"].between(
        ts["净负荷预测下界_kWh"], ts["净负荷预测上界_kWh"], inclusive="both"
    )
    ts["上界突破"] = ts["实际净负荷_kWh"] > ts["净负荷预测上界_kWh"]
    ts["下界突破"] = ts["实际净负荷_kWh"] < ts["净负荷预测下界_kWh"]
    calibration = (
        ts.groupby("区间编号", as_index=False)
        .agg(
            经验覆盖率=("区间覆盖", "mean"),
            上界突破率=("上界突破", "mean"),
            下界突破率=("下界突破", "mean"),
            平均区间宽度_kWh=("净负荷预测上界_kWh", lambda x: 0.0),
        )
    )
    width = (
        ts.assign(宽度=ts["净负荷预测上界_kWh"] - ts["净负荷预测下界_kWh"])
        .groupby("区间编号")["宽度"].mean()
    )
    calibration["平均区间宽度_kWh"] = calibration["区间编号"].map(width)
    return metric_frame, calibration


def add_check(
    rows: list[dict[str, object]],
    category: str,
    item: str,
    value: float | int | str,
    criterion: str,
    passed: bool,
) -> None:
    rows.append(
        {"类别": category, "检查项": item, "实际值": value, "通过标准": criterion, "是否通过": bool(passed)}
    )


def validate_model_outputs(
    ts: pd.DataFrame,
    daily: pd.DataFrame,
    events: pd.DataFrame,
    workbook_path: Path,
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    add_check(rows, "数据", "日期数量", ts["日期"].nunique(), "334", ts["日期"].nunique() == EXPECTED_DAYS)
    add_check(rows, "数据", "10分钟记录数", len(ts), str(EXPECTED_DAYS * T), len(ts) == EXPECTED_DAYS * T)
    per_day_counts = ts.groupby("日期")["区间编号"].nunique()
    add_check(rows, "数据", "每天区间完整", int((per_day_counts == T).sum()), "334天均为144点", bool((per_day_counts == T).all()))
    add_check(rows, "数据", "区间编号范围", f"{ts['区间编号'].min()}-{ts['区间编号'].max()}", "1-144", ts["区间编号"].min() == 1 and ts["区间编号"].max() == 144)

    plan_load_kwh = ts["预测负载_kW"].to_numpy() / 6.0 + ts["计划风险增量_kWh"].to_numpy()
    plan_balance = (
        ts["计划购电量_kWh"].to_numpy()
        + ts["预测光伏_kW"].to_numpy() / 6.0
        + ts["计划放电量_kWh"].to_numpy()
        - plan_load_kwh
        - ts["计划充电量_kWh"].to_numpy()
        - ts["计划剩余电量_kWh"].to_numpy()
    )
    actual_balance = (
        ts["计划购电量_kWh"].to_numpy()
        + ts["紧急购电量_kWh"].to_numpy()
        + ts["实际光伏_kWh"].to_numpy()
        + ts["实际放电量_kWh"].to_numpy()
        - ts["实际负载_kWh"].to_numpy()
        - ts["实际充电量_kWh"].to_numpy()
        - ts["总剩余电量_kWh"].to_numpy()
    )
    soc_residual = ts["实际SOC_kWh"].to_numpy() - (
        ts["实际期初SOC_kWh"].to_numpy()
        + ETA_C * ts["实际充电量_kWh"].to_numpy()
        - ts["实际放电量_kWh"].to_numpy() / ETA_D
    )
    checks = [
        ("计划平衡最大残差_kWh", float(np.max(np.abs(plan_balance))), 1.0e-4),
        ("实际平衡最大残差_kWh", float(np.max(np.abs(actual_balance))), 1.0e-4),
        ("SOC递推最大残差_kWh", float(np.max(np.abs(soc_residual))), 1.0e-4),
    ]
    for item, value, tolerance in checks:
        add_check(rows, "约束", item, value, f"<={tolerance}", value <= tolerance)
    min_soc, max_soc = float(ts["实际SOC_kWh"].min()), float(ts["实际SOC_kWh"].max())
    add_check(rows, "约束", "SOC最小值_kWh", min_soc, f">={E_MIN}", min_soc >= E_MIN - 1.0e-5)
    add_check(rows, "约束", "SOC最大值_kWh", max_soc, f"<={E_MAX}", max_soc <= E_MAX + 1.0e-5)
    max_charge = float(ts["实际充电量_kWh"].max())
    max_discharge = float(ts["实际放电量_kWh"].max())
    add_check(rows, "约束", "单时段最大充电量_kWh", max_charge, f"<={STEP_MAX_KWH:.6f}", max_charge <= STEP_MAX_KWH + 1.0e-5)
    add_check(rows, "约束", "单时段最大放电量_kWh", max_discharge, f"<={STEP_MAX_KWH:.6f}", max_discharge <= STEP_MAX_KWH + 1.0e-5)
    simultaneous = int(((ts["实际充电量_kWh"] > 1.0e-6) & (ts["实际放电量_kWh"] > 1.0e-6)).sum())
    add_check(rows, "约束", "同时充放电时段数", simultaneous, "0", simultaneous == 0)
    negative_columns = [
        "计划购电量_kWh", "实际充电量_kWh", "实际放电量_kWh", "紧急购电量_kWh",
        "总剩余电量_kWh", "弃光量_kWh",
    ]
    negative_count = int((ts[negative_columns] < -1.0e-8).sum().sum())
    add_check(rows, "约束", "负决策值数量", negative_count, "0", negative_count == 0)

    coverage = float(
        (
            (ts["实际净负荷_kWh"] >= ts["净负荷预测下界_kWh"])
            & (ts["实际净负荷_kWh"] <= ts["净负荷预测上界_kWh"])
        ).mean()
    )
    add_check(rows, "风险", "净负荷区间经验覆盖率", coverage, ">=0.85", coverage >= 0.85)

    day_state = ts.groupby("日期").agg(期初=("实际期初SOC_kWh", "first"), 期末=("实际SOC_kWh", "last"))
    continuity = day_state["期初"].iloc[1:].to_numpy() - day_state["期末"].iloc[:-1].to_numpy()
    continuity_max = float(np.max(np.abs(continuity))) if len(continuity) else 0.0
    add_check(rows, "约束", "跨日SOC连续最大残差_kWh", continuity_max, "<=0.0001", continuity_max <= 1.0e-4)

    recomputed = (
        ts.assign(
            计划费=ts["电价_元每kWh"] * ts["计划购电量_kWh"],
            紧急费=5.0 * ts["电价_元每kWh"] * ts["紧急购电量_kWh"],
        )
        .groupby("日期", as_index=False)
        .agg(计划费=("计划费", "sum"), 紧急费=("紧急费", "sum"))
    )
    compare = daily.merge(recomputed, on="日期", validate="one_to_one")
    plan_cost_error = float(np.max(np.abs(compare["计划购电费_元"] - compare["计划费"])))
    emergency_cost_error = float(np.max(np.abs(compare["紧急购电费_元"] - compare["紧急费"])))
    add_check(rows, "费用", "全天计划费最大复算误差_元", plan_cost_error, "<=0.01", plan_cost_error <= 0.01)
    add_check(rows, "费用", "每日紧急费最大复算误差_元", emergency_cost_error, "<=0.01", emergency_cost_error <= 0.01)
    event_sum = float(pd.to_numeric(events["购电量"], errors="coerce").fillna(0).sum())
    detail_sum = float(ts["紧急购电量_kWh"].sum())
    add_check(rows, "费用", "紧急事件与明细电量误差_kWh", abs(event_sum - detail_sum), "<=0.001", abs(event_sum - detail_sum) <= 0.001)

    if not workbook_path.exists():
        add_check(rows, "工作簿", "result2存在", "不存在", "存在", False)
        return pd.DataFrame(rows)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    expected_sheets = ["计划购电量", "充放电量", "紧急购电量"]
    add_check(rows, "工作簿", "工作表名称", ",".join(workbook.sheetnames), ",".join(expected_sheets), workbook.sheetnames == expected_sheets)
    ws_plan = workbook["计划购电量"]
    ws_storage = workbook["充放电量"]
    ws_emergency = workbook["紧急购电量"]
    add_check(rows, "工作簿", "计划表数据行数", ws_plan.max_row - 1, "334", ws_plan.max_row - 1 == EXPECTED_DAYS)
    add_check(rows, "工作簿", "计划表列数", ws_plan.max_column, "147", ws_plan.max_column == 147)
    add_check(rows, "工作簿", "充放电表数据行数", ws_storage.max_row - 1, str(EXPECTED_DAYS * 6), ws_storage.max_row - 1 == EXPECTED_DAYS * 6)
    ellipsis_count = 0
    for ws in (ws_storage, ws_emergency):
        for row in ws.iter_rows(values_only=True):
            ellipsis_count += sum(value in {"……", "…", "⁝"} for value in row)
    add_check(rows, "工作簿", "残留省略号数量", ellipsis_count, "0", ellipsis_count == 0)

    # 全量核对计划表144列的循环映射。
    workbook_plan = np.array(
        [
            [0.0 if value is None else float(value) for value in row[1:145]]
            for row in ws_plan.iter_rows(
                min_row=2, max_row=ws_plan.max_row, min_col=1, max_col=145,
                values_only=True,
            )
        ],
        dtype=float,
    )
    detail_plan = ts.pivot(index="日期", columns="区间编号", values="计划购电量_kWh").sort_index().to_numpy()
    expected_rotated = np.concatenate([detail_plan[:, 1:], detail_plan[:, :1]], axis=1)
    mapping_error = float(np.max(np.abs(workbook_plan - expected_rotated)))
    add_check(rows, "工作簿", "计划量循环映射最大误差_kWh", mapping_error, "<=0.001", mapping_error <= 0.001)
    return pd.DataFrame(rows)


def monthly_summary(ts: pd.DataFrame, daily: pd.DataFrame) -> pd.DataFrame:
    frame = daily.copy()
    frame["月份"] = frame["日期"].dt.month
    monthly = (
        frame.groupby("月份", as_index=False)
        .agg(
            天数=("日期", "size"),
            计划购电费_元=("计划购电费_元", "sum"),
            紧急购电费_元=("紧急购电费_元", "sum"),
            紧急购电量_kWh=("紧急购电量_kWh", "sum"),
            弃光量_kWh=("弃光量_kWh", "sum"),
            未利用外网电量_kWh=("未利用外网电量_kWh", "sum"),
            平均负载MAE_kW=("预测负载MAE_kW", "mean"),
            平均光伏MAE_kW=("预测光伏MAE_kW", "mean"),
        )
    )
    pv_month = (
        ts.assign(月份=ts["日期"].dt.month)
        .groupby("月份")["实际光伏_kWh"].sum()
    )
    monthly["弃光率"] = monthly["弃光量_kWh"] / monthly["月份"].map(pv_month)
    monthly["总购电费_元"] = monthly["计划购电费_元"] + monthly["紧急购电费_元"]
    return monthly


def plot_daily_forecast_error(daily: pd.DataFrame, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(15, 6))
    ax.plot(daily["日期"], daily["预测负载MAE_kW"], label="负载MAE", color="#2563EB", lw=1.3)
    ax.plot(daily["日期"], daily["预测光伏MAE_kW"], label="光伏MAE", color="#F59E0B", lw=1.3)
    ax.set_title("滚动预测每日误差")
    ax.set_xlabel("日期")
    ax.set_ylabel("MAE (kW)")
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m月"))
    ax.legend()
    ax.grid(True, linestyle="--", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_key_date_net_load(ts: pd.DataFrame, path: Path) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(15, 10), sharex=True)
    for ax, date in zip(axes.ravel(), KEY_DATES):
        part = ts[ts["日期"].eq(date)].sort_values("区间编号")
        x = part["区间编号"]
        actual = part["实际净负荷_kWh"] * 6.0
        forecast = part["预测净负荷_kWh"] * 6.0
        lower = part["净负荷预测下界_kWh"] * 6.0
        upper = part["净负荷预测上界_kWh"] * 6.0
        ax.fill_between(x, lower, upper, color="#93C5FD", alpha=0.35, label="预测风险区间")
        ax.plot(x, forecast, color="#2563EB", lw=1.5, label="预测净负荷")
        ax.plot(x, actual, color="#111827", lw=1.6, label="实际净负荷")
        ax.axhline(0, color="#64748B", lw=0.8)
        ax.set_title(date.strftime("%Y-%m-%d"))
        ax.set_ylabel("功率 (kW)")
        ax.grid(True, linestyle="--", alpha=0.25)
    for ax in axes[-1]:
        ax.set_xticks([1, 37, 73, 109, 144])
        ax.set_xticklabels(["0:10", "6:10", "12:10", "18:10", "24:00"])
        ax.set_xlabel("右端点时刻")
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.legend(handles, labels, ncol=3, loc="upper center")
    fig.suptitle("指定日期净负荷预测区间与实际值", y=0.995, fontsize=16)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_key_date_operation(ts: pd.DataFrame, path: Path) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(15, 10), sharex=True)
    for ax, date in zip(axes.ravel(), KEY_DATES):
        part = ts[ts["日期"].eq(date)].sort_values("区间编号")
        x = part["区间编号"]
        ax.plot(x, part["实际SOC_kWh"], color="#7C3AED", lw=1.8, label="SOC")
        ax.set_ylabel("SOC (kWh)")
        ax.set_title(date.strftime("%Y-%m-%d"))
        ax2 = ax.twinx()
        ax2.bar(x, part["紧急购电量_kWh"], width=1.0, color="#DC2626", alpha=0.60, label="紧急购电")
        ax2.plot(x, part["实际充电量_kWh"], color="#16A34A", lw=0.9, alpha=0.8, label="充电")
        ax2.plot(x, -part["实际放电量_kWh"], color="#F59E0B", lw=0.9, alpha=0.8, label="放电（负向显示）")
        ax2.set_ylabel("时段电量 (kWh)")
        ax.grid(True, linestyle="--", alpha=0.25)
    for ax in axes[-1]:
        ax.set_xticks([1, 37, 73, 109, 144])
        ax.set_xticklabels(["0:10", "6:10", "12:10", "18:10", "24:00"])
        ax.set_xlabel("右端点时刻")
    # 用固定代理线，避免额外twinx对象干扰图面。
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    proxies = [
        Line2D([0], [0], color="#7C3AED", lw=2),
        Line2D([0], [0], color="#16A34A", lw=2),
        Line2D([0], [0], color="#F59E0B", lw=2),
        Patch(facecolor="#DC2626", alpha=0.60),
    ]
    fig.legend(proxies, ["SOC", "充电", "放电", "紧急购电"], ncol=4, loc="upper center")
    fig.suptitle("指定日期储能运行与紧急购电", y=0.995, fontsize=16)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_daily_cost(daily: pd.DataFrame, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(15, 6))
    ax.plot(daily["日期"], daily["计划购电费_元"] / 1000, color="#2563EB", lw=1.3, label="计划购电费")
    ax.fill_between(daily["日期"], 0, daily["紧急购电费_元"] / 1000, color="#DC2626", alpha=0.45, label="紧急购电费")
    ax.set_title("每日计划购电费与紧急购电费")
    ax.set_xlabel("日期")
    ax.set_ylabel("费用（千元）")
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m月"))
    ax.legend()
    ax.grid(True, linestyle="--", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_heatmap(ts: pd.DataFrame, value: str, title: str, cbar_label: str, cmap: str, path: Path) -> None:
    matrix = ts.pivot(index="日期", columns="区间编号", values=value).sort_index()
    fig, ax = plt.subplots(figsize=(16, 8))
    values = matrix.to_numpy()
    positive_values = values[values > 0]
    vmax = float(np.quantile(positive_values, 0.98)) if len(positive_values) else 1.0
    sns.heatmap(matrix, cmap=cmap, vmin=0, vmax=vmax, xticklabels=False, yticklabels=False,
                cbar_kws={"label": cbar_label}, ax=ax)
    first_dates = matrix.reset_index().assign(月份=matrix.index.month).groupby("月份", as_index=False).first()["日期"]
    positions = [matrix.index.get_loc(date) + 0.5 for date in first_dates]
    ax.set_yticks(positions)
    ax.set_yticklabels([f"{date.month}月" for date in first_dates], rotation=0)
    ax.set_xticks(np.arange(0, T, 18) + 0.5)
    ax.set_xticklabels(["0:10", "3:10", "6:10", "9:10", "12:10", "15:10", "18:10", "21:10"])
    ax.set_title(title)
    ax.set_xlabel("右端点时刻")
    ax.set_ylabel("日期")
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_calibration(calibration: pd.DataFrame, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(15, 6))
    x = calibration["区间编号"]
    ax.plot(x, calibration["经验覆盖率"], color="#2563EB", lw=1.8, label="经验覆盖率")
    ax.plot(x, calibration["上界突破率"], color="#DC2626", lw=1.2, label="上界突破率")
    ax.plot(x, calibration["下界突破率"], color="#F59E0B", lw=1.2, label="下界突破率")
    ax.set_ylim(0, 1.02)
    ax.set_xticks([1, 37, 73, 109, 144])
    ax.set_xticklabels(["0:10", "6:10", "12:10", "18:10", "24:00"])
    ax.set_title("净负荷预测风险区间的经验覆盖")
    ax.set_xlabel("右端点时刻")
    ax.set_ylabel("比例")
    ax.legend(ncol=3)
    ax.grid(True, linestyle="--", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def plot_monthly_tradeoff(monthly: pd.DataFrame, path: Path) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(15, 6))
    axes[0].bar(monthly["月份"], monthly["计划购电费_元"] / 1e6, color="#2563EB", label="计划费")
    axes[0].bar(monthly["月份"], monthly["紧急购电费_元"] / 1e6,
                bottom=monthly["计划购电费_元"] / 1e6, color="#DC2626", label="紧急费")
    axes[0].set_title("月度购电费用构成")
    axes[0].set_xlabel("月份")
    axes[0].set_ylabel("费用（百万元）")
    axes[0].legend()
    scatter = axes[1].scatter(
        monthly["弃光率"] * 100,
        monthly["紧急购电量_kWh"] / 1000,
        c=monthly["月份"], cmap="viridis", s=90,
    )
    for _, row in monthly.iterrows():
        axes[1].annotate(f"{int(row['月份'])}月", (row["弃光率"] * 100, row["紧急购电量_kWh"] / 1000),
                         xytext=(4, 4), textcoords="offset points", fontsize=9)
    axes[1].set_title("月度弃光与紧急购电")
    axes[1].set_xlabel("弃光率（%）")
    axes[1].set_ylabel("紧急购电量（MWh）")
    axes[1].grid(True, linestyle="--", alpha=0.3)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="问题2结果检验、诊断和可视化")
    parser.add_argument(
        "--input-dir", type=Path, default=SCRIPT_DIR / "q2_output"
    )
    parser.add_argument("--result2", type=Path, default=None, help="默认读取input-dir/result2_filled.xlsx")
    parser.add_argument(
        "--output-dir", type=Path, default=SCRIPT_DIR / "q2_diagnostics"
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    configure_plot_style()
    result2_path = args.result2 or args.input_dir / "result2_filled.xlsx"
    args.output_dir.mkdir(parents=True, exist_ok=True)
    ts, daily, events = load_outputs(args.input_dir)
    forecast_metrics, calibration = forecast_diagnostics(ts)
    validation = validate_model_outputs(ts, daily, events, result2_path)
    monthly = monthly_summary(ts, daily)

    forecast_metrics.to_csv(args.output_dir / "q2_forecast_metrics.csv", index=False, encoding="utf-8-sig")
    calibration.to_csv(args.output_dir / "q2_interval_calibration.csv", index=False, encoding="utf-8-sig")
    validation.to_csv(args.output_dir / "q2_validation_report.csv", index=False, encoding="utf-8-sig")
    monthly.to_csv(args.output_dir / "q2_monthly_summary.csv", index=False, encoding="utf-8-sig")

    plot_daily_forecast_error(daily, args.output_dir / "01_每日预测误差.png")
    plot_key_date_net_load(ts, args.output_dir / "02_指定日期净负荷预测区间.png")
    plot_key_date_operation(ts, args.output_dir / "03_指定日期储能与紧急购电.png")
    plot_daily_cost(daily, args.output_dir / "04_每日费用构成.png")
    plot_heatmap(ts, "紧急购电量_kWh", "全年紧急购电热力图", "紧急购电量 (kWh)", "Reds", args.output_dir / "05_紧急购电热力图.png")
    plot_heatmap(ts, "弃光量_kWh", "全年弃光热力图", "弃光量 (kWh)", "YlOrBr", args.output_dir / "06_弃光热力图.png")
    plot_calibration(calibration, args.output_dir / "07_预测区间覆盖率.png")
    plot_monthly_tradeoff(monthly, args.output_dir / "08_月度费用风险与弃光.png")

    failed = validation.loc[~validation["是否通过"]]
    print("\n========== 预测指标 ==========")
    print(forecast_metrics.round(4).to_string(index=False))
    print(f"\n全时段经验区间覆盖率：{((ts['实际净负荷_kWh'] >= ts['净负荷预测下界_kWh']) & (ts['实际净负荷_kWh'] <= ts['净负荷预测上界_kWh'])).mean():.2%}")
    print("\n========== 约束、费用和工作簿检验 ==========")
    print(validation.to_string(index=False))
    if failed.empty:
        print("\n[通过] 所有自动检验均通过。")
    else:
        print(f"\n[警告] 有{len(failed)}项检验未通过，请查看q2_validation_report.csv。")
    print(f"[完成] 诊断结果：{args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
