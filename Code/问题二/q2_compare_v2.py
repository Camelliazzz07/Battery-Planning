#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""新版独立回测/诊断：运行四个方案，比较费用、期末库存、备用和预测。

pip install numpy pandas scipy openpyxl matplotlib seaborn
python q2_compare_v2.py --attachment1 附件1.xlsx --attachment2 附件2.xlsx --template result2.xlsx
已有完整输出时：python q2_compare_v2.py --skip-run --runs-dir q2_v2_runs

依赖同目录q2_run_model_v2.py。所有方案共同用旧策略预热1月，不重置SOC。
比较是固定参数的历史回测，不是未知数据上的效果保证；不根据全年结果自动调参。
输出费用对照CSV/MD、约束检查、预测覆盖率及4张英文图（避免中文字体依赖）。
"""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import load_workbook
import q2_run_model_v2 as model

VARIANTS = ["baseline", "forecast_only", "soc_only", "revised"]
LABELS = {"baseline": "Original", "forecast_only": "Forecast only",
          "soc_only": "SOC only", "revised": "Revised"}


def checks(ts, daily, events, path, variant):
    records = []
    def check(name, error, tolerance=1e-4):
        value = float(error)
        records.append({"方案": variant, "检查项": name, "误差或违规数": value,
                        "容差": tolerance, "通过": bool(np.isfinite(value) and value <= tolerance)})
    def absmax(x):
        return np.max(np.abs(np.asarray(x)))
    def violation(x):
        return max(0.0, float(np.max(np.asarray(x))))
    check("334天且日期完整", int(not pd.DatetimeIndex(daily["日期"]).equals(
        pd.date_range("2025-02-01", "2025-12-31"))), 0)
    check("48096个10分钟点", abs(len(ts)-334*144), 0)
    check("每一天144点", absmax(ts.groupby("日期").size()-144), 0)
    check("序号自然日排序", absmax(ts["区间编号"].to_numpy()-np.tile(np.arange(1,145),334)), 0)
    for prefix in ["计划", "实际"]:
        c, d, e = [ts[prefix+x+"_kWh"].to_numpy() for x in ["充电量", "放电量", "SOC"]]
        if prefix == "实际":
            start = ts["实际期初SOC_kWh"].to_numpy()
        else:
            em = e.reshape(-1,144)
            start = np.column_stack([daily["期初SOC_kWh"].to_numpy(), em[:,:-1]]).ravel()
        check(prefix+"SOC递推", absmax(e-start-0.9*c+d/0.9))
        check(prefix+"SOC下限", violation(1200-e))
        check(prefix+"SOC上限", violation(e-10800))
        check(prefix+"充电功率", violation(c-5000/6))
        check(prefix+"放电功率", violation(d-5000/6))
        check(prefix+"充放电互斥", int(((c>1e-5)&(d>1e-5)).sum()), 0)
        check(prefix+"充放电非负", max(violation(-c), violation(-d)))
    q = ts["计划购电量_kWh"].to_numpy()
    emergency = ts["紧急购电量_kWh"].to_numpy()
    price = ts["电价_元每kWh"].to_numpy()
    check("计划量非负", violation(-q))
    check("紧急量非负", violation(-emergency))
    check("实际供需平衡", absmax(q+emergency+ts["实际光伏_kWh"]+ts["实际放电量_kWh"]
          -ts["实际负载_kWh"]-ts["实际充电量_kWh"]-ts["总剩余电量_kWh"]))
    check("计划供需平衡含剩余", absmax(q+ts["预测光伏_kW"]/6+ts["计划放电量_kWh"]
          -ts["预测负载_kW"]/6-ts["计划风险增量_kWh"]-ts["计划充电量_kWh"]-ts["计划剩余电量_kWh"]))
    check("跨日SOC连续", absmax(daily["期初SOC_kWh"].to_numpy()[1:]-daily["期末SOC_kWh"].to_numpy()[:-1]))
    check("计划费按计划量结算", abs(np.dot(q,price)-daily["计划购电费_元"].sum()))
    check("紧急费按5倍电价", abs(np.dot(emergency,5*price)-daily["紧急购电费_元"].sum()))
    check("事件汇总与10分钟紧急量一致", abs(events["购电量"].sum()-emergency.sum()))
    if variant == "revised":
        e = ts["计划SOC_kWh"].to_numpy().reshape(-1,144)
        start = np.column_stack([daily["期初SOC_kWh"].to_numpy(),e[:,:-1]]).ravel()
        up = ts["正向累计备用_kWh"].to_numpy()/0.9
        down = ts["反向累计备用_kWh"].to_numpy()*0.9
        su, sd = ts["正向备用缺口_kWh"].to_numpy(), ts["反向备用缺口_kWh"].to_numpy()
        check("开始SOC上备用含软缺口", violation(1200+up-start-su))
        check("开始SOC下备用含软缺口", violation(start-sd-10800+down))
        check("结束SOC上备用含软缺口", violation(1200+up-e.ravel()-su))
        check("结束SOC下备用含软缺口", violation(e.ravel()-sd-10800+down))
        check("放电功率备用含软缺口", violation(ts["计划放电量_kWh"]+ts["正向瞬时备用_kWh"]
              -ts["正向功率备用缺口_kWh"]-5000/6))
        check("充电功率备用含软缺口", violation(ts["计划充电量_kWh"]+ts["反向瞬时备用_kWh"]
              -ts["反向功率备用缺口_kWh"]-5000/6))
    wb = load_workbook(path, read_only=True, data_only=True)
    check("result2工作表", int(wb.sheetnames != ["计划购电量","充放电量","紧急购电量"]),0)
    plan = list(wb["计划购电量"].iter_rows(values_only=True))
    check("计划表334天", abs(len(plan)-335),0)
    headers = [str(h).replace("7:0-", "7:00-") for h in plan[0][1:145]]
    if headers[-1] == "0:00-0:10+1":
        headers[-1] = "0:00+1-0:10+1"
    expected = [model.natural_interval_label(i) for i in range(1,145)]
    check("模板全部时段标签", sum(x!=y for x,y in zip(headers, expected)),0)
    qm = q.reshape(-1,144)
    expected_q = np.concatenate([qm[:,1:],qm[:,:1]],axis=1)
    check("144列首尾映射", absmax(np.asarray([row[1:145] for row in plan[1:]],float)-expected_q))
    check("模板每日日期", sum(pd.Timestamp(row[0])!=day for row,day in zip(plan[1:],daily["日期"])),0)
    check("模板计划总电量", absmax(np.array([row[145] for row in plan[1:]])-qm.sum(axis=1)))
    check("模板计划费用", absmax(np.array([row[146] for row in plan[1:]])-daily["计划购电费_元"].to_numpy()))
    storage = list(wb["充放电量"].iter_rows(values_only=True))[1:]
    check("充放电表2004行", abs(len(storage)-334*6),0)
    cm=ts["实际充电量_kWh"].to_numpy().reshape(-1,24).sum(axis=1)
    dm=ts["实际放电量_kWh"].to_numpy().reshape(-1,24).sum(axis=1)
    check("四小时累计充电", absmax(np.array([r[2] for r in storage])-cm))
    check("四小时累计放电", absmax(np.array([r[3] for r in storage])-dm))
    check("模板0点SOC", absmax(np.array([r[5] for r in storage[::6]])-daily["期初SOC_kWh"].to_numpy()))
    check("模板24点SOC", absmax(np.array([r[5] for r in storage[1::6]])-daily["期末SOC_kWh"].to_numpy()))
    event_rows = list(wb["紧急购电量"].iter_rows(values_only=True))[1:]
    check("模板紧急量", abs(sum(float(r[2] or 0) for r in event_rows)-emergency.sum()))
    wb.close()
    return records


def logic_tests():
    """与数据表现无关的单元测试：预测不偷看未来，备用按剩余误差计算。"""
    rng=np.random.default_rng(20260912)
    a=rng.uniform(100,200,(40,144)); b=a.copy(); b[30:]=1e9
    for lead in [0,1]:
        f,_=model.seasonal_load_forecast(a,30,a[0],lead)
        g,_=model.seasonal_load_forecast(b,30,b[0],lead)
        np.testing.assert_array_equal(f,g)
    hist=[np.full(144,100.0) for _ in range(10)]
    cfg=model.ModelConfig()
    r=model.revised_risk_profiles(hist,np.ones(144)*500,np.zeros(144),cfg)
    np.testing.assert_allclose(r["plan_margin"],100)
    np.testing.assert_allclose(r["up_cumulative"],0)
    np.testing.assert_allclose(r["down_cumulative"],0)
    rng_hist=[rng.normal(0,100,144) for _ in range(20)]
    r=model.revised_risk_profiles(rng_hist,np.ones(144),np.zeros(144),cfg)
    errors=np.array(rng_hist)
    np.testing.assert_allclose(r["interval_upper"],np.quantile(errors,.95,axis=0))
    np.testing.assert_allclose(r["interval_lower"],np.quantile(errors,.05,axis=0))


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--attachment1",type=Path,default=Path("附件1.xlsx"))
    ap.add_argument("--attachment2",type=Path,default=Path("附件2.xlsx"))
    ap.add_argument("--template",type=Path,default=Path("result2.xlsx"))
    ap.add_argument("--runs-dir",type=Path,default=Path("q2_v2_runs"))
    ap.add_argument("--output-dir",type=Path,default=Path("q2_v2_comparison"))
    ap.add_argument("--skip-run",action="store_true")
    args=ap.parse_args()
    args.output_dir.mkdir(parents=True,exist_ok=True)
    if not args.skip_run:
        for variant in VARIANTS:
            cmd=[sys.executable,str(Path(__file__).with_name("q2_run_model_v2.py")),
                 "--attachment1",str(args.attachment1),"--attachment2",str(args.attachment2),
                 "--template",str(args.template),"--variant",variant,
                 "--output-dir",str(args.runs_dir/variant)]
            subprocess.run(cmd,check=True)
    logic_tests()
    rows, validations, monthly_rows = [], [], []
    all_daily={}
    for variant in VARIANTS:
        folder=args.runs_dir/variant
        ts=pd.read_csv(folder/"q2_timeseries.csv",parse_dates=["日期"])
        daily=pd.read_csv(folder/"q2_daily_summary.csv",parse_dates=["日期"])
        events=pd.read_csv(folder/"q2_emergency_events.csv")
        meta=json.loads((folder/"q2_run_metadata.json").read_text(encoding="utf-8"))
        if meta["config"]["variant"] != variant:
            raise ValueError(f"{folder}方案元数据不一致")
        validations.extend(checks(ts,daily,events,folder/"result2_filled.xlsx",variant))
        covered=ts["实际净负荷_kWh"].between(ts["净负荷预测下界_kWh"],ts["净负荷预测上界_kWh"])
        rows.append({
            "方案":variant,
            "计划购电费_元":daily["计划购电费_元"].sum(),
            "紧急购电费_元":daily["紧急购电费_元"].sum(),
            "总购电费_元":daily["总购电费_元"].sum(),
            "紧急购电量_kWh":daily["紧急购电量_kWh"].sum(),
            "紧急购电天数":int((daily["紧急购电量_kWh"]>1e-6).sum()),
            "总剩余电量_kWh":ts["总剩余电量_kWh"].sum(),
            "归因弃光率":ts["弃光量_kWh"].sum()/ts["实际光伏_kWh"].sum(),
            "负载MAE_kW":daily["预测负载MAE_kW"].mean(),
            "预测区间经验覆盖率":covered.mean(),
            "实际期初SOC_kWh":daily.iloc[0]["期初SOC_kWh"],
            "实际期末SOC_kWh":daily.iloc[-1]["期末SOC_kWh"],
            "计划日末SOC最小_kWh":daily["计划日末SOC_kWh"].min(),
            "计划日末SOC最大_kWh":daily["计划日末SOC_kWh"].max(),
            "上能量备用软缺口时段数":int((ts["正向备用缺口_kWh"]>1e-5).sum()),
            "上功率备用软缺口时段数":int((ts["正向功率备用缺口_kWh"]>1e-5).sum()),
        })
        m=daily.groupby(daily["日期"].dt.to_period("M"))["总购电费_元"].sum()
        monthly_rows.extend({"方案":variant,"月份":str(k),"总购电费_元":v} for k,v in m.items())
        all_daily[variant]=daily
    result=pd.DataFrame(rows)
    if np.ptp(result["实际期初SOC_kWh"].to_numpy())>1e-5:
        raise ValueError("期初库存不一致，对照不公平")
    pmin=float(ts["电价_元每kWh"].min()); pmax=float(ts["电价_元每kWh"].max())
    base=result.iloc[0]
    result["相对原模型节约_元"]=base["总购电费_元"]-result["总购电费_元"]
    result["相对原模型节约率"]=result["相对原模型节约_元"]/base["总购电费_元"]
    # 库存调整只是比较指标，不修改真实购电费，不伪造售电收入。
    result["库存调整后比较费_元"]=result["总购电费_元"]-0.9*pmin*(
        result["实际期末SOC_kWh"]-result["实际期初SOC_kWh"])
    result["保守扣除期末差异后节约_元"]=result["相对原模型节约_元"]-0.9*5*pmax*np.abs(
        result["实际期末SOC_kWh"]-base["实际期末SOC_kWh"])
    report=pd.DataFrame(validations)
    report.to_csv(args.output_dir/"q2_validation_v2.csv",index=False,encoding="utf-8-sig")
    if not report["通过"].all():
        print(report.loc[~report["通过"]].to_string(index=False))
        raise RuntimeError("存在物理或结果表检查失败，不能交付为通过结果")
    result.to_csv(args.output_dir/"q2_cost_comparison.csv",index=False,encoding="utf-8-sig")
    pd.DataFrame(monthly_rows).to_csv(args.output_dir/"q2_monthly_comparison.csv",index=False,encoding="utf-8-sig")
    sns.set_theme(style="whitegrid",context="notebook")
    colors=sns.color_palette("colorblind",4)
    names=[LABELS[v] for v in VARIANTS]
    fig,ax=plt.subplots(figsize=(9,5))
    x=np.arange(4)
    ax.bar(x,result["计划购电费_元"]/1e6,label="Planned purchase",color="#4C78A8")
    ax.bar(x,result["紧急购电费_元"]/1e6,bottom=result["计划购电费_元"]/1e6,label="Emergency purchase",color="#E45756")
    ax.set_xticks(x,names); ax.set_ylabel("Cost (million CNY)");ax.set_title("Billed cost: 2025-02-01 to 2025-12-31")
    for i,v in enumerate(result["总购电费_元"]): ax.text(i,v/1e6+.05,f"{v/1e6:.3f}",ha="center")
    ax.legend();fig.tight_layout();fig.savefig(args.output_dir/"01_cost_comparison.png",dpi=160);plt.close(fig)
    fig,ax=plt.subplots(figsize=(11,4))
    for i,variant in enumerate(VARIANTS[1:],1):
        saving=all_daily["baseline"]["总购电费_元"]-all_daily[variant]["总购电费_元"]
        ax.plot(all_daily[variant]["日期"],saving.cumsum()/1e4,label=LABELS[variant],color=colors[i])
    ax.axhline(0,color="gray",lw=.7);ax.set_ylabel("Cumulative saving (10,000 CNY)");ax.legend();ax.set_title("Savings against original model")
    fig.tight_layout();fig.savefig(args.output_dir/"02_cumulative_savings.png",dpi=160);plt.close(fig)
    fig,axs=plt.subplots(2,1,figsize=(11,7),sharex=True)
    for i,variant in enumerate(["baseline","revised"]):
        d=all_daily[variant]
        axs[0].plot(d["日期"],d["计划日末SOC_kWh"],label=LABELS[variant],alpha=.8)
        axs[1].plot(d["日期"],d["期末SOC_kWh"],label=LABELS[variant],alpha=.8)
    for ax,label in zip(axs,["Planned day-end energy (kWh)","Actual day-end energy (kWh)"]):
        ax.axhline(6000,color="gray",ls="--",lw=.8,label="6000 reference (not v2 constraint)")
        ax.set_ylabel(label);ax.legend(fontsize=8)
    fig.tight_layout();fig.savefig(args.output_dir/"03_day_end_soc.png",dpi=160);plt.close(fig)
    fig,axs=plt.subplots(1,2,figsize=(11,4))
    axs[0].bar(names,result["紧急购电量_kWh"]/1e3,color=colors);axs[0].set_ylabel("Emergency energy (MWh)")
    axs[1].bar(names,result["总剩余电量_kWh"]/1e3,color=colors);axs[1].set_ylabel("Total surplus energy (MWh)")
    for ax in axs: ax.tick_params(axis="x",rotation=20)
    fig.tight_layout();fig.savefig(args.output_dir/"04_emergency_and_surplus.png",dpi=160);plt.close(fig)
    rev=result.loc[result["方案"]=="revised"].iloc[0]
    text=["# 问题二新版对照结果", "", "评价期：2025-02-01至2025-12-31，共334天。",
          "各方案1月统一旧策略预热，2月1日使用同一期初库存；无跨日重置。",
          "", "| 方案 | 计划费（元） | 紧急费（元） | 合计（元） | 期末电量（kWh） |",
          "|---|---:|---:|---:|---:|"]
    for _,r in result.iterrows():
        text.append(f"| {r['方案']} | {r['计划购电费_元']:,.2f} | {r['紧急购电费_元']:,.2f} | {r['总购电费_元']:,.2f} | {r['实际期末SOC_kWh']:,.2f} |")
    text.extend(["",f"完整版相对原模型节约 {rev['相对原模型节约_元']:,.2f} 元（{rev['相对原模型节约率']:.2%}）。",
        f"即使用5倍最高电价乘放电效率对期末电量差作保守扣减，节约仍为 {rev['保守扣除期末差异后节约_元']:,.2f} 元。此为敏感性度量，不是实际补购方案。",
        f"约束、时间映射与费用复算共 {len(report)} 项通过。",
        "", "## 改动及限制", "",
        "- SOC：48h滚动计划，只执行当天；每天6000目标与罚金均移除，窗口末使用保守电量残值。",
        "- 预测：负载取上周同一天（不是近7天平均）；光伏保留原集成。切换时按历史滚动预测重算校准残差。",
        "- 备用：先扣当前加购量，再算累计残差最大前缀分位数；跨午夜使用相邻历史日。移除平滑和人为截断，软缺口明确输出。",
        "- 同时检查开始库存与计划后库存；功率备用和能量备用分别约束；备用缺口不代表负荷停供，仍由紧急购电保障。",
        "- 第二天PV延续当次预测、备用复用首日曲线，是48h展望近似。日内仍为因果计划跟踪，不是每10分钟MPC。",
        "- 仍是带风险备用的代理目标，不含紧急费用情景期望或CVaR；求解最优不等于实际总费用全局最优。",
        "- 0.8/0.95分位数、4元上备用罚金、0.02元下备用罚金等是公开默认假设，未用全期表现自动挑选。",
        "- 新版预测区间Q05-Q95名义覆盖90%，不能与旧版被平滑截断的区间简单等同，也不承诺95%整日可靠。",
        "- 总剩余电量为稳定指标；弃光量按优先弃光记账，不能唯一识别来源。",
        "- 原result2末列表头为0:00-0:10+1；保留该原文，按既定约定填本日第1点。7:0视为7:00。",
        "- 本次为已知数据的固定方案回测；不是独立未见数据测试，也不是保证未来省钱。",
        "", "## 运行", "", "```bash", "python -m pip install numpy pandas scipy openpyxl matplotlib seaborn",
        "python q2_run_model_v2.py --output-dir q2_v2_runs/revised",
        "python q2_compare_v2.py", "```", "",
        "主脚本单独即可填result2；比较脚本运行四方案并独立检验。不需要旧主脚本。",
        "已有四方案完整输出时，比较脚本可加 --skip-run。"])
    (args.output_dir/"q2_v2_comparison.md").write_text("\n".join(text)+"\n",encoding="utf-8")
    print(result.to_string(index=False))
    print(f"\n检查通过：{len(report)} 项。结果目录：{args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
