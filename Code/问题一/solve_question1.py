# -*- coding: utf-8 -*-
"""2026 C题问题1：微网全天计划购电与储能充放电优化。

依赖：
    pip install pandas numpy scipy openpyxl

在 VS Code 中可直接运行当前文件。默认从仓库“附件”目录读取附件1.xlsx，
从“附件/附件5”读取result1.xlsx，结果仍保存到本脚本所在的“问题一”目录：
    python solve_question1.py

指定路径：
    python solve_question1.py --data 附件1.xlsx --template result1.xlsx \
        --output result1_问题1求解结果.xlsx

建模约定：
1. 附件1采用右端点记时，0:10 表示 0:00-0:10；不插值、不补行。
2. 功率乘 dt=1/6 h 后转换为每段电量（kWh）。
3. 计划购电量 x_t 直接进入费用目标，费用按计划量结算。
4. 供电约束使用 >=，允许光伏或供电富余被弃用。
5. c_t 是交流侧充电电量，d_t 是储能向交流侧交付的放电电量。
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix, vstack


for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8", errors="replace")


T = 144
DT = 1.0 / 6.0
ETA_CHARGE = 0.90
ETA_DISCHARGE = 0.90
SOC_INITIAL = 6000.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
POWER_MAX_KW = 5000.0
ENERGY_MAX_KWH = POWER_MAX_KW * DT

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent
ATTACHMENT_DIR = PROJECT_ROOT / "附件"
TEMPLATE_DIR = ATTACHMENT_DIR / "附件5"

REQUIRED_COLUMNS = ["时间", "电价", "小区负载", "光伏发电预测功率"]


def parse_time_to_minutes(value: object) -> int:
    """解析附件时间；0:00+1 返回1440，其余返回小时*60+分钟。"""
    if pd.isna(value):
        raise ValueError("时间列存在空值。")
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.hour * 60 + value.minute
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float, np.integer, np.floating)):
        number = float(value)
        if 0 <= number <= 1:
            return int(round(number * 1440))
        raise ValueError(f"无法识别数值型时间：{value!r}")

    text = str(value).strip().replace("：", ":")
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(\+1)?", text)
    if not match:
        raise ValueError(f"无法解析时间：{value!r}")
    hour = int(match.group(1))
    minute = int(match.group(2))
    next_day = match.group(3) is not None
    if hour > 23 or minute > 59:
        raise ValueError(f"时间超出范围：{value!r}")
    if next_day:
        if hour != 0 or minute != 0:
            raise ValueError(f"只允许0:00+1表示次日零点：{value!r}")
        return 1440
    return hour * 60 + minute


def read_data(path: Path) -> pd.DataFrame:
    """读取、校验并按分钟升序排列附件1。时间轴异常时停止求解。"""
    df = pd.read_excel(path, engine="openpyxl")
    df.columns = df.columns.astype(str).str.strip()
    missing = [column for column in REQUIRED_COLUMNS if column not in df.columns]
    if missing:
        raise KeyError(f"附件1缺少列：{missing}")

    df = df[REQUIRED_COLUMNS].copy()
    df["分钟数"] = df["时间"].map(parse_time_to_minutes)
    df = df.sort_values("分钟数", kind="stable").reset_index(drop=True)

    if len(df) != T:
        raise ValueError(f"附件1必须有{T}行，实际为{len(df)}行。")
    expected_minutes = np.arange(10, 1441, 10)
    actual_minutes = df["分钟数"].to_numpy(dtype=int)
    if not np.array_equal(actual_minutes, expected_minutes):
        raise ValueError(
            "时间轴必须严格为10,20,...,1440分钟；请检查缺失或重复时刻。"
        )

    for column in REQUIRED_COLUMNS[1:]:
        df[column] = pd.to_numeric(df[column], errors="coerce")
        if df[column].isna().any():
            bad_rows = (df.index[df[column].isna()] + 2).tolist()
            raise ValueError(
                f"列“{column}”存在缺失或非数值数据，Excel行号：{bad_rows}。"
                "优化脚本不自动插值，请先完成数据清洗。"
            )
    if (df["电价"] < 0).any():
        raise ValueError("问题1模型未设置购电上限，不允许出现负电价。")

    df["负载电量_kWh"] = df["小区负载"] * DT
    df["光伏电量_kWh"] = df["光伏发电预测功率"] * DT
    return df


def variable_slices() -> dict[str, slice]:
    """变量顺序：[购电x, 充电c, 放电d, SOC E(0..144), 状态z]。"""
    return {
        "x": slice(0, T),
        "charge": slice(T, 2 * T),
        "discharge": slice(2 * T, 3 * T),
        "soc": slice(3 * T, 4 * T + 1),
        "mode": slice(4 * T + 1, 5 * T + 1),
    }


def solve_model(df: pd.DataFrame) -> dict[str, np.ndarray | float]:
    """建立并求解MILP。"""
    idx = variable_slices()
    n_vars = 5 * T + 1

    objective = np.zeros(n_vars)
    objective[idx["x"]] = df["电价"].to_numpy(dtype=float)

    lower = np.zeros(n_vars)
    upper = np.full(n_vars, np.inf)
    lower[idx["soc"]] = SOC_MIN
    upper[idx["soc"]] = SOC_MAX
    upper[idx["charge"]] = ENERGY_MAX_KWH
    upper[idx["discharge"]] = ENERGY_MAX_KWH
    upper[idx["mode"]] = 1.0

    # 固定0:00和24:00的储电量为6000 kWh。
    soc_indices = np.arange(idx["soc"].start, idx["soc"].stop)
    lower[soc_indices[0]] = upper[soc_indices[0]] = SOC_INITIAL
    lower[soc_indices[-1]] = upper[soc_indices[-1]] = SOC_INITIAL

    constraints: list[LinearConstraint] = []

    # 1) 区间电量平衡：x_t + PV_t + d_t >= L_t + c_t。
    balance = lil_matrix((T, n_vars))
    balance[np.arange(T), np.arange(idx["x"].start, idx["x"].stop)] = 1.0
    balance[np.arange(T), np.arange(idx["charge"].start, idx["charge"].stop)] = -1.0
    balance[np.arange(T), np.arange(idx["discharge"].start, idx["discharge"].stop)] = 1.0
    net_demand = (
        df["负载电量_kWh"].to_numpy(dtype=float)
        - df["光伏电量_kWh"].to_numpy(dtype=float)
    )
    constraints.append(LinearConstraint(balance.tocsr(), net_demand, np.inf))

    # 2) SOC递推：E_t-E_(t-1)-eta_c*c_t+d_t/eta_d=0。
    dynamics = lil_matrix((T, n_vars))
    for t in range(T):
        dynamics[t, idx["soc"].start + t + 1] = 1.0
        dynamics[t, idx["soc"].start + t] = -1.0
        dynamics[t, idx["charge"].start + t] = -ETA_CHARGE
        dynamics[t, idx["discharge"].start + t] = 1.0 / ETA_DISCHARGE
    constraints.append(LinearConstraint(dynamics.tocsr(), 0.0, 0.0))

    # 3) 禁止同时充放电：c_t <= M*z_t；d_t <= M*(1-z_t)。
    charge_mode = lil_matrix((T, n_vars))
    discharge_mode = lil_matrix((T, n_vars))
    for t in range(T):
        charge_mode[t, idx["charge"].start + t] = 1.0
        charge_mode[t, idx["mode"].start + t] = -ENERGY_MAX_KWH
        discharge_mode[t, idx["discharge"].start + t] = 1.0
        discharge_mode[t, idx["mode"].start + t] = ENERGY_MAX_KWH
    constraints.append(LinearConstraint(charge_mode.tocsr(), -np.inf, 0.0))
    constraints.append(
        LinearConstraint(discharge_mode.tocsr(), -np.inf, ENERGY_MAX_KWH)
    )

    integrality = np.zeros(n_vars, dtype=int)
    integrality[idx["mode"]] = 1

    result = milp(
        c=objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=constraints,
        options={"disp": False, "mip_rel_gap": 1e-9},
    )
    if not result.success or result.x is None:
        raise RuntimeError(f"优化失败：status={result.status}, message={result.message}")

    solution = result.x
    purchase = np.maximum(solution[idx["x"]], 0.0)
    charge = np.maximum(solution[idx["charge"]], 0.0)
    discharge = np.maximum(solution[idx["discharge"]], 0.0)
    soc = solution[idx["soc"]]

    # 小数误差清理，不改变有意义的最优解。
    for array in (purchase, charge, discharge):
        array[np.abs(array) < 1e-7] = 0.0

    return {
        "purchase": purchase,
        "charge": charge,
        "discharge": discharge,
        "soc": soc,
        "objective": float(np.dot(df["电价"].to_numpy(), purchase)),
    }


def validate_solution(df: pd.DataFrame, solution: dict[str, np.ndarray | float]) -> None:
    """独立检查电量平衡、SOC递推、容量和终端约束。"""
    x = np.asarray(solution["purchase"])
    charge = np.asarray(solution["charge"])
    discharge = np.asarray(solution["discharge"])
    soc = np.asarray(solution["soc"])
    load = df["负载电量_kWh"].to_numpy()
    pv = df["光伏电量_kWh"].to_numpy()

    balance_slack = x + pv + discharge - load - charge
    dynamics_error = soc[1:] - soc[:-1] - ETA_CHARGE * charge + discharge / ETA_DISCHARGE

    tolerance = 1e-5
    checks = {
        "供电不低于负载": float(balance_slack.min()) >= -tolerance,
        "SOC递推": float(np.max(np.abs(dynamics_error))) <= tolerance,
        "SOC安全下限": float(soc.min()) >= SOC_MIN - tolerance,
        "SOC安全上限": float(soc.max()) <= SOC_MAX + tolerance,
        "初始SOC": abs(float(soc[0]) - SOC_INITIAL) <= tolerance,
        "终止SOC": abs(float(soc[-1]) - SOC_INITIAL) <= tolerance,
        "不同时充放电": not np.any((charge > tolerance) & (discharge > tolerance)),
    }
    failed = [name for name, passed in checks.items() if not passed]
    if failed:
        raise RuntimeError(f"求解结果未通过检查：{failed}")
    print("约束检查：全部通过。")


def fill_result_workbook(
    template_path: Path,
    output_path: Path,
    solution: dict[str, np.ndarray | float],
) -> None:
    """保留模板结构并填入计划购电量、6段充放电量及首尾SOC。"""
    workbook = load_workbook(template_path)
    required_sheets = ["计划购电量", "充放电量"]
    missing_sheets = [name for name in required_sheets if name not in workbook.sheetnames]
    if missing_sheets:
        raise KeyError(f"result1模板缺少工作表：{missing_sheets}")

    purchase_sheet = workbook["计划购电量"]
    storage_sheet = workbook["充放电量"]
    purchase = np.asarray(solution["purchase"])
    charge = np.asarray(solution["charge"])
    discharge = np.asarray(solution["discharge"])
    soc = np.asarray(solution["soc"])

    # 官方模板标签从0:10-0:20开始，因此按[第2段,...,第144段,第1段]循环移位。
    template_purchase = np.concatenate([purchase[1:], purchase[:1]])
    expected_first = "0:10-0:20"
    expected_last = "0:00+1-0:10+1"
    if str(purchase_sheet["A2"].value).strip() != expected_first:
        raise ValueError(f"模板A2应为{expected_first}，实际为{purchase_sheet['A2'].value!r}")
    if str(purchase_sheet["A145"].value).strip() != expected_last:
        raise ValueError(f"模板A145应为{expected_last}，实际为{purchase_sheet['A145'].value!r}")

    for row, value in enumerate(template_purchase, start=2):
        purchase_sheet.cell(row=row, column=2, value=round(float(value), 6))

    # 每4小时含24个10分钟区间，共6段。
    charge_blocks = charge.reshape(6, 24).sum(axis=1)
    discharge_blocks = discharge.reshape(6, 24).sum(axis=1)
    for block in range(6):
        storage_sheet.cell(row=block + 2, column=2, value=round(float(charge_blocks[block]), 6))
        storage_sheet.cell(row=block + 2, column=3, value=round(float(discharge_blocks[block]), 6))

    storage_sheet["E2"] = round(float(soc[0]), 6)
    storage_sheet["E3"] = round(float(soc[-1]), 6)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def print_summary(df: pd.DataFrame, solution: dict[str, np.ndarray | float]) -> None:
    purchase = np.asarray(solution["purchase"])
    charge = np.asarray(solution["charge"])
    discharge = np.asarray(solution["discharge"])
    soc = np.asarray(solution["soc"])
    print("\n========== 问题1求解摘要 ==========")
    print(f"全天计划购电量：{purchase.sum():.6f} kWh")
    print(f"全天计划购电费：{float(solution['objective']):.6f} 元")
    print(f"全天充电量：{charge.sum():.6f} kWh")
    print(f"全天放电量：{discharge.sum():.6f} kWh")
    print(f"SOC范围：{soc.min():.6f} - {soc.max():.6f} kWh")
    print(f"0:00 SOC：{soc[0]:.6f} kWh")
    print(f"24:00 SOC：{soc[-1]:.6f} kWh")

    # 论文表1指定时段按自然时间轴取值，而不是按模板行号直接索引。
    specified_starts = ["10:00", "12:00", "14:00", "16:00", "18:00", "20:00"]
    print("\n论文表1指定时段：")
    for start in specified_starts:
        hour, minute = map(int, start.split(":"))
        natural_index = (hour * 60 + minute) // 10
        end_minutes = hour * 60 + minute + 10
        end_hour, end_minute = divmod(end_minutes, 60)
        print(
            f"{start}-{end_hour:02d}:{end_minute:02d}："
            f"{purchase[natural_index]:.6f} kWh"
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="求解C题问题1并填充result1模板")
    parser.add_argument(
        "--data", type=Path, default=ATTACHMENT_DIR / "附件1.xlsx"
    )
    parser.add_argument(
        "--template", type=Path, default=TEMPLATE_DIR / "result1.xlsx"
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=SCRIPT_DIR / "result1_问题1求解结果.xlsx",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    for path, label in [(args.data, "附件1"), (args.template, "result1模板")]:
        if not path.exists():
            raise FileNotFoundError(f"找不到{label}：{path.resolve()}")

    df = read_data(args.data)
    print("数据检查：144个右端点时段完整，功率已转换为区间电量。")
    solution = solve_model(df)
    validate_solution(df, solution)
    print_summary(df, solution)
    fill_result_workbook(args.template, args.output, solution)
    print(f"\n结果文件已生成：{args.output.resolve()}")


if __name__ == "__main__":
    main()
