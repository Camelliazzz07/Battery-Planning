# -*- coding: utf-8 -*-
"""问题3：因果预测更新、风险备用、滚动优化与实际运行回测。

在 VS Code 中可直接运行。默认读取仓库“附件”目录，正式结果保存到
“提交结果/result3.xlsx”，其余明细保存到“问题三/辅助输出”。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
import scipy
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix

for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8", errors="replace")

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent
ATTACHMENT_DIR = PROJECT_ROOT / "附件"
TEMPLATE_DIR = ATTACHMENT_DIR / "附件5"
RESULT_DIR = PROJECT_ROOT / "提交结果"
AUXILIARY_DIR = SCRIPT_DIR / "辅助输出"

T, DT = 144, 1 / 6
E_MIN, E_MAX, E_INIT, P_MAX = 1200.0, 10800.0, 6000.0, 5000.0
ISSUES = (0, 6, 12, 18)
SELECTED = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")


@dataclass(frozen=True)
class Config:
    settlement: str = "no_refund"
    eta_c: float = 0.9
    eta_d: float = 0.9
    plan_quantile: float = 0.8
    update_quantile: float = 0.7
    reserve_quantile: float = 0.95
    reserve_horizon: int = 24  # 4 hours, in ten-minute intervals
    lookback_days: int = 56
    min_residual_days: int = 7
    reserve_weight: float = 1.0
    reserve_penalty: float = 4.0
    terminal_target: float = 6000.0
    terminal_penalty: float = 0.65
    load_bias_weight: float = 0.0  # lag-7 primary baseline; optional causal update
    mip_gap: float = 1e-5
    time_limit: float = 30.0
    spill_tiebreak: float = 1e-7


@dataclass
class Data:
    dates: pd.DatetimeIndex
    price: np.ndarray
    prior_load: np.ndarray
    load: np.ndarray
    pv: np.ndarray
    forecasts: dict[tuple[pd.Timestamp, int], np.ndarray]


def endpoint_minute(value) -> int:
    if isinstance(value, (datetime, time, pd.Timestamp)):
        return value.hour * 60 + value.minute or 1440
    if isinstance(value, (float, int, np.floating, np.integer)):
        if 0 <= value <= 1:
            return int(round(float(value) * 1440)) or 1440
    s = str(value).strip().replace("：", ":")
    if s in ("0:00+1", "00:00+1", "24:00"):
        return 1440
    m = re.fullmatch(r"(\d{1,2}):(\d{2})", s)
    if not m or int(m[1]) > 23 or int(m[2]) > 59:
        raise ValueError(f"Invalid time: {value!r}")
    return int(m[1]) * 60 + int(m[2]) or 1440


def clock(minute: int) -> str:
    return f"{minute // 60}:{minute % 60:02d}"


def interval(t: int) -> str:
    return f"{clock(10*t)}-{clock(10*(t+1))}"


def finite_nonnegative(a, name):
    a = np.asarray(a, dtype=float)
    if not np.isfinite(a).all() or (a < 0).any():
        raise ValueError(f"{name}: missing/nonfinite/negative values")
    return a


def read_inputs(root: Path, end: str = "2025-12-31") -> Data:
    a1 = pd.read_excel(root / "附件1.xlsx")
    minutes = np.array([endpoint_minute(x) for x in a1["时间"]])
    order = np.argsort(minutes)
    if not np.array_equal(minutes[order], np.arange(10, 1441, 10)):
        raise ValueError("附件1 must contain all 144 right endpoints")
    price = finite_nonnegative(a1["电价"].to_numpy()[order], "price")
    prior = finite_nonnegative(a1["小区负载"].to_numpy()[order], "prior_load")
    arrays, dates = [], None
    for sheet in ("小区负载", "光伏发电实际功率"):
        frame = pd.read_excel(root / "附件2.xlsx", sheet_name=sheet)
        current = pd.DatetimeIndex(pd.to_datetime(frame.iloc[:, 0])).normalize()
        if dates is not None and not current.equals(dates):
            raise ValueError("Load/PV dates differ")
        dates = current
        mins = np.array([endpoint_minute(c) for c in frame.columns[1:]])
        ix = np.argsort(mins)
        if not np.array_equal(mins[ix], np.arange(10, 1441, 10)):
            raise ValueError(f"{sheet}: missing/duplicate endpoint")
        arrays.append(finite_nonnegative(frame.iloc[:, 1:].to_numpy()[:, ix], sheet))
    if not dates.equals(pd.date_range("2025-01-01", dates[-1])):
        raise ValueError("Require contiguous dates starting 2025-01-01 for SOC warm-up")
    if pd.Timestamp(end) not in dates or pd.Timestamp(end) < pd.Timestamp("2025-02-01"):
        raise ValueError("end must be a covered date on/after 2025-02-01")
    keep = dates <= pd.Timestamp(end)
    dates = dates[keep]
    a3 = pd.read_excel(root / "附件3.xlsx")
    a3["日期"] = pd.to_datetime(a3["日期"].ffill()).dt.normalize()
    forecasts = {}
    lead_cols = [f"预报{i}小时" for i in range(1, 25)]
    for _, row in a3.iterrows():
        minute = endpoint_minute(row["预报时刻"]) % 1440
        if minute % 60 or minute // 60 not in ISSUES:
            raise ValueError(f'Unexpected forecast issue: {row["预报时刻"]}')
        key = (row["日期"], minute // 60)
        if key in forecasts:
            raise ValueError(f"Duplicate forecast issue: {key}")
        forecasts[key] = finite_nonnegative(row[lead_cols].to_numpy(), str(key))
    for day in dates:
        for hour in ISSUES:
            if (day, hour) not in forecasts:
                raise ValueError(f"Missing forecast: {day} {hour}:00")
    return Data(dates, price, prior, arrays[0][keep], arrays[1][keep], forecasts)


def forecast_at(data: Data, day_index: int, hour: int, cfg: Config):
    """Information contract: actuals < day + hour; only THIS issued PV row.

    Attachment 3 lead j targets issue_timestamp+j hours, including next day.
    This day-only optimizer consumes just the remaining same-day portion.
    """
    s = hour * 6
    if day_index >= 7:
        load = data.load[day_index - 7].copy()
    elif day_index:
        load = data.load[:day_index].mean(axis=0)
    else:
        load = data.prior_load.copy()
    if s and cfg.load_bias_weight:
        # Optional observed last-hour mean error, decaying into future.
        bias = np.mean(data.load[day_index, s - 6 : s] - load[s - 6 : s])
        load[s:] += cfg.load_bias_weight * bias * np.exp(-np.arange(T - s) / 36)
    load = np.maximum(load[s:], 0.0)
    # Boundary observation is known at the release instant (no look-ahead).
    anchor = (
        data.pv[day_index, s - 1]
        if s
        else data.pv[day_index - 1, -1] if day_index else 0.0
    )
    published = data.forecasts[(data.dates[day_index], hour)]
    x = np.arange(25) * 60 + hour * 60
    targets = np.arange(s + 1, T + 1) * 10
    pv = np.interp(targets, x, np.r_[anchor, published])
    return load * DT, pv * DT


def make_forecast_cache(data: Data, cfg: Config):
    """Precompute causal forecasts; residuals are sliced to days STRICTLY before use."""
    cache, errors = {}, {}
    for h in ISSUES:
        pairs = [forecast_at(data, i, h, cfg) for i in range(len(data.dates))]
        cache[h] = (np.array([p[0] for p in pairs]), np.array([p[1] for p in pairs]))
        predicted = cache[h][0] - cache[h][1]
        errors[h] = (data.load[:, h * 6 :] - data.pv[:, h * 6 :]) * DT - predicted
    return cache, errors


def risk_profile(
    history: np.ndarray, load: np.ndarray, pv: np.ndarray, quantile: float, cfg: Config
):
    """Margin first; reserve quantile on remaining error, no quantile subtraction."""
    n = len(load)
    if len(history) < cfg.min_residual_days:
        # Jan cold-start only; an explicit engineering assumption, not data fit.
        margin = 0.05 * load + 0.08 * pv
        return margin, np.zeros(n)
    residuals = history[-cfg.lookback_days :]
    margin = np.maximum(np.quantile(residuals, quantile, axis=0), 0.0)
    remaining = residuals - margin[None, :]
    reserve = np.zeros(n)
    for t in range(n):
        paths = np.cumsum(remaining[:, t : min(n, t + cfg.reserve_horizon)], axis=1)
        maxima = np.maximum(paths.max(axis=1), 0.0)
        reserve[t] = cfg.reserve_weight * np.quantile(maxima, cfg.reserve_quantile)
    # Reserve is grid-side kWh at START of interval. No undocumented clipping.
    return margin, reserve


def settlement(price, q0, active, emergency, mode):
    """One final-reference settlement per delivery interval; never sum snapshots."""
    up = np.maximum(active - q0, 0)
    down = np.maximum(q0 - active, 0)
    return {
        "up_kwh": up,
        "down_kwh": down,
        "plan_cost": price * q0,
        "up_cost": 1.5 * price * up,
        "down_penalty": 0.5 * price * down,
        "refund": price * down if mode == "refund" else np.zeros_like(down),
        "emergency_cost": 5 * price * emergency,
    }


def solve_remaining(price, load, pv, start_soc, margin, reserve, cfg: Config, q0=None):
    """Deterministic proxy MILP; q0=None for midnight, otherwise final-reference cost."""
    n = len(price)
    groups = ("q", "c", "d", "e", "w", "slack", "z", "u", "v")
    off = {g: i * n for i, g in enumerate(groups)}
    pos, neg = len(groups) * n, len(groups) * n + 1
    nv = neg + 1
    obj, lo, hi = np.zeros(nv), np.zeros(nv), np.full(nv, np.inf)
    ix = lambda name: slice(off[name], off[name] + n)
    max_step = P_MAX * DT
    hi[ix("c")] = hi[ix("d")] = max_step
    lo[ix("e")], hi[ix("e")] = E_MIN, E_MAX
    hi[ix("z")] = 1
    obj[ix("w")] = cfg.spill_tiebreak
    obj[ix("slack")] = cfg.reserve_penalty
    obj[pos] = obj[neg] = cfg.terminal_penalty
    if q0 is None:
        obj[ix("q")] = price
        hi[ix("u")] = hi[ix("v")] = 0
    else:
        obj[ix("u")] = 1.5 * price
        obj[ix("v")] = (0.5 if cfg.settlement == "no_refund" else -0.5) * price
        hi[ix("v")] = q0  # no negative deliveries
        if cfg.settlement == "no_refund":
            # Free spill: q<q0 is dominated. Added purchases may still be canceled.
            hi[ix("v")] = 0
            lo[ix("q")] = q0
    integer = np.zeros(nv)
    integer[ix("z")] = 1
    neq = 2 * n + 1 + (n if q0 is not None else 0)
    aeq, beq = lil_matrix((neq, nv)), np.zeros(neq)
    aub, bub = lil_matrix((3 * n, nv)), np.zeros(3 * n)
    for t in range(n):
        # q + PV + d = load + margin + c + spill
        for g, coefficient in [("q", 1), ("d", 1), ("c", -1), ("w", -1)]:
            aeq[t, off[g] + t] = coefficient
        beq[t] = load[t] + margin[t] - pv[t]
        aeq[n + t, off["e"] + t] = 1
        aeq[n + t, off["c"] + t] = -cfg.eta_c
        aeq[n + t, off["d"] + t] = 1 / cfg.eta_d
        if t:
            aeq[n + t, off["e"] + t - 1] = -1
        else:
            beq[n + t] = start_soc
        aub[t, off["c"] + t], aub[t, off["z"] + t] = 1, -max_step
        aub[n + t, off["d"] + t], aub[n + t, off["z"] + t] = 1, max_step
        bub[n + t] = max_step
        # E_start + slack >= Emin + reserve / eta_d.
        aub[2 * n + t, off["slack"] + t] = -1
        if t:
            aub[2 * n + t, off["e"] + t - 1] = -1
            bub[2 * n + t] = -E_MIN - reserve[t] / cfg.eta_d
        else:
            bub[2 * n + t] = start_soc - E_MIN - reserve[t] / cfg.eta_d
        if q0 is not None:
            row = 2 * n + 1 + t
            for g, coefficient in [("q", 1), ("u", -1), ("v", 1)]:
                aeq[row, off[g] + t] = coefficient
            beq[row] = q0[t]
    aeq[2 * n, off["e"] + n - 1], aeq[2 * n, pos], aeq[2 * n, neg] = 1, -1, 1
    beq[2 * n] = cfg.terminal_target
    sol = milp(
        obj,
        integrality=integer,
        bounds=Bounds(lo, hi),
        constraints=[
            LinearConstraint(aeq.tocsr(), beq, beq),
            LinearConstraint(aub.tocsr(), -np.inf, bub),
        ],
        options={
            "mip_rel_gap": cfg.mip_gap,
            "time_limit": cfg.time_limit,
            "presolve": True,
        },
    )
    if not sol.success or sol.x is None:
        raise RuntimeError(f"MILP failed: {sol.status} {sol.message}")
    equality_error = np.max(np.abs(aeq @ sol.x - beq))
    inequality_error = max(0, float(np.max(aub @ sol.x - bub)))
    if max(equality_error, inequality_error) > 2e-4:
        raise RuntimeError("MILP primal feasibility audit failed")
    out = {g: sol.x[ix(g)].copy() for g in groups}
    out["q"] = np.maximum(out["q"], 0)
    if q0 is not None and cfg.settlement == "no_refund":
        out["q"] = np.maximum(out["q"], q0)
    out.update(
        objective=float(sol.fun),
        mip_gap=float(sol.mip_gap),
        mip_absolute_gap=float(abs(sol.fun - sol.mip_dual_bound)),
        max_constraint_error=float(max(equality_error, inequality_error)),
    )
    return out


def dispatch(q, load, pv, e, cfg: Config):
    """Within-interval balancing feedback: use ONLY current net power and SOC.

    Assumes the 10-minute sampled power represents this interval and fast feedback
    is available. No future realized load/PV enters this rule.
    """
    surplus = q + pv - load
    c = min(max(surplus, 0.0), P_MAX * DT, max(0, (E_MAX - e) / cfg.eta_c))
    d = min(max(-surplus, 0.0), P_MAX * DT, max(0, (e - E_MIN) * cfg.eta_d))
    spill = max(0.0, surplus - c)
    emergency = max(0.0, -surplus - d)
    return c, d, e + cfg.eta_c * c - d / cfg.eta_d, emergency, spill


def simulate_day(data, cache, errors, di, e, hours, cfg):
    q0, active = None, np.zeros(T)
    records, revisions, solves = [], [], []
    for t in range(T):
        h = t // 6
        if t % 6 == 0 and h in hours:
            load, pv = cache[h][0][di], cache[h][1][di]
            hist = errors[h][max(0, di - cfg.lookback_days) : di]
            margin, reserve = risk_profile(
                hist,
                load,
                pv,
                cfg.plan_quantile if h == 0 else cfg.update_quantile,
                cfg,
            )
            old = active[t:].copy()
            sol = solve_remaining(
                data.price[t:],
                load,
                pv,
                e,
                margin,
                reserve,
                cfg,
                None if h == 0 else q0[t:],
            )
            if h == 0:
                q0 = sol["q"].copy()
            active[t:] = sol["q"]  # preceding intervals cannot be overwritten
            solves.append(
                {
                    "date": str(data.dates[di].date()),
                    "issue_hour": h,
                    "start_soc_kwh": e,
                    "objective_proxy": sol["objective"],
                    "mip_gap": sol["mip_gap"],
                    "mip_absolute_gap": sol["mip_absolute_gap"],
                    "constraint_error": sol["max_constraint_error"],
                    "max_reserve_slack_kwh": float(sol["slack"].max()),
                    "historical_days": len(hist),
                }
            )
            for j in range(len(load)):
                revisions.append(
                    {
                        "date": str(data.dates[di].date()),
                        "issue_hour": h,
                        "interval": interval(t + j),
                        "t": t + j,
                        "forecast_load_kwh": load[j],
                        "forecast_pv_kwh": pv[j],
                        "margin_kwh": margin[j],
                        "reserve_grid_kwh": reserve[j],
                        "q0_kwh": q0[t + j],
                        "previous_active_kwh": old[j] if h else q0[j],
                        "new_active_kwh": sol["q"][j],
                        "planned_charge_kwh": sol["c"][j],
                        "planned_discharge_kwh": sol["d"][j],
                        "planned_end_soc_kwh": sol["e"][j],
                        "reserve_slack_soc_kwh": sol["slack"][j],
                    }
                )
        start_e = e
        load, pv = data.load[di, t] * DT, data.pv[di, t] * DT
        c, d, e, r, w = dispatch(active[t], load, pv, e, cfg)
        records.append(
            {
                "date": str(data.dates[di].date()),
                "t": t,
                "start_time": str(data.dates[di] + pd.Timedelta(minutes=t * 10)),
                "end_time": str(data.dates[di] + pd.Timedelta(minutes=(t + 1) * 10)),
                "interval": interval(t),
                "price": data.price[t],
                "q0_kwh": q0[t],
                "active_kwh": active[t],
                "load_kwh": load,
                "pv_kwh": pv,
                "charge_kwh": c,
                "discharge_kwh": d,
                "start_soc_kwh": start_e,
                "end_soc_kwh": e,
                "emergency_kwh": r,
                "spill_kwh": w,
            }
        )
    df = pd.DataFrame(records)
    fees = settlement(
        df.price.to_numpy(),
        df.q0_kwh.to_numpy(),
        df.active_kwh.to_numpy(),
        df.emergency_kwh.to_numpy(),
        cfg.settlement,
    )
    for key, value in fees.items():
        df[key] = value
    df["total_cost"] = (
        df.plan_cost + df.up_cost + df.down_penalty - df.refund + df.emergency_cost
    )
    df["storage_loss_kwh"] = (1 - cfg.eta_c) * df.charge_kwh + (
        1 / cfg.eta_d - 1
    ) * df.discharge_kwh
    return e, df, revisions, solves


def audit(df, cfg):
    # Independent actual power balance and chronological SOC checks.
    balance = (
        df.active_kwh
        + df.emergency_kwh
        + df.pv_kwh
        + df.discharge_kwh
        - df.load_kwh
        - df.charge_kwh
        - df.spill_kwh
    )
    recursion = (
        df.end_soc_kwh
        - df.start_soc_kwh
        - cfg.eta_c * df.charge_kwh
        + df.discharge_kwh / cfg.eta_d
    )
    continuity = df.start_soc_kwh.to_numpy()[1:] - df.end_soc_kwh.to_numpy()[:-1]
    qdiff = df.active_kwh - df.q0_kwh
    # Alternative fee formula, independent of split fee columns.
    coeff = 0.5 if cfg.settlement == "no_refund" else -0.5
    reconstructed = df.price * (
        df.q0_kwh
        + 1.5 * qdiff.clip(lower=0)
        + coeff * (-qdiff).clip(lower=0)
        + 5 * df.emergency_kwh
    )
    result = {
        "max_balance_error_kwh": float(abs(balance).max()),
        "max_soc_recursion_error_kwh": float(abs(recursion).max()),
        "max_continuity_error_kwh": float(np.max(abs(continuity), initial=0)),
        "max_settlement_error_yuan": float(abs(reconstructed - df.total_cost).max()),
        "min_soc_kwh": float(min(df.start_soc_kwh.min(), df.end_soc_kwh.min())),
        "max_soc_kwh": float(max(df.start_soc_kwh.max(), df.end_soc_kwh.max())),
        "max_charge_kw": float(df.charge_kwh.max() / DT),
        "max_discharge_kw": float(df.discharge_kwh.max() / DT),
        "simultaneous_charge_discharge_count": int(
            ((df.charge_kwh > 1e-6) & (df.discharge_kwh > 1e-6)).sum()
        ),
        "nonfinite_numeric_count": int(
            (~np.isfinite(df.select_dtypes("number"))).to_numpy().sum()
        ),
    }
    expected_start = pd.to_datetime(df.date) + pd.to_timedelta(df.t * 10, unit="min")
    assert (pd.to_datetime(df.start_time) == expected_start).all()
    assert (
        pd.to_datetime(df.end_time) - pd.to_datetime(df.start_time)
        == pd.Timedelta(minutes=10)
    ).all()
    assert (
        df.groupby("date")
        .t.apply(lambda x: np.array_equal(x.to_numpy(), np.arange(T)))
        .all()
    )
    for key in (
        "max_balance_error_kwh",
        "max_soc_recursion_error_kwh",
        "max_continuity_error_kwh",
        "max_settlement_error_yuan",
    ):
        assert result[key] < 2e-4, (key, result[key])
    assert (
        E_MIN - 1e-5 <= result["min_soc_kwh"] <= result["max_soc_kwh"] <= E_MAX + 1e-5
    )
    assert max(result["max_charge_kw"], result["max_discharge_kw"]) <= P_MAX + 1e-5
    assert (
        result["simultaneous_charge_discharge_count"]
        == result["nonfinite_numeric_count"]
        == 0
    )
    if cfg.settlement == "no_refund":
        assert (df.active_kwh >= df.q0_kwh - 1e-6).all()
    result["passed"] = True
    return result


def daily_summary(df):
    additive = [
        "q0_kwh",
        "active_kwh",
        "up_kwh",
        "down_kwh",
        "emergency_kwh",
        "charge_kwh",
        "discharge_kwh",
        "spill_kwh",
        "storage_loss_kwh",
        "plan_cost",
        "up_cost",
        "down_penalty",
        "refund",
        "emergency_cost",
        "total_cost",
    ]
    days = df.groupby("date")[additive].sum()
    days["start_soc_kwh"] = df.groupby("date").start_soc_kwh.first()
    days["end_soc_kwh"] = df.groupby("date").end_soc_kwh.last()
    return days.reset_index()


def emergency_events(df):
    rows = []
    for date, day in df.groupby("date", sort=False):
        values = day.emergency_kwh.to_numpy()
        t, found = 0, False
        while t < T:
            if values[t] <= 1e-6:
                t += 1
                continue
            a = t
            while t + 1 < T and values[t + 1] > 1e-6:
                t += 1
            rows.append(
                [
                    date,
                    f"{clock(10*a)}-{clock(10*(t+1))}",
                    float(values[a : t + 1].sum()),
                ]
            )
            found, t = True, t + 1
        if not found:
            rows.append([date, None, None])
    return pd.DataFrame(rows, columns=["日期", "购电时间段", "购电量"])


def table_payload(df, daily, events):
    """Natural chronological columns: explicitly correct the template's shifted labels."""
    labels = [interval(t) for t in range(T)]
    plan_rows, active_rows, storage_rows = [], [], []
    for date, day in df.groupby("date", sort=False):
        plan_rows.append(
            [date]
            + day.q0_kwh.tolist()
            + [float(day.q0_kwh.sum()), float(day.plan_cost.sum())]
        )
        active_rows.append(
            [date]
            + day.active_kwh.tolist()
            + [float(day.active_kwh.sum()), float(day.total_cost.sum())]
        )
        for b in range(6):
            block = day.iloc[b * 24 : (b + 1) * 24]
            storage_rows.append(
                [
                    date if b == 0 else None,
                    f"{4*b}:00-{4*(b+1)}:00",
                    float(block.charge_kwh.sum()),
                    float(block.discharge_kwh.sum()),
                    "0:00" if b == 0 else "24:00" if b == 1 else None,
                    (
                        float(day.start_soc_kwh.iloc[0])
                        if b == 0
                        else float(day.end_soc_kwh.iloc[-1]) if b == 1 else None
                    ),
                ]
            )
    headers = ["日期\\时间"] + labels + ["全天购电量", "全天购电费"]
    return {
        "计划购电量": {"headers": headers, "rows": plan_rows},
        "调整购电量": {"headers": headers, "rows": active_rows},
        "充放电量": {
            "headers": ["日期", "时间段", "充电量", "放电量", "时刻", "储电量"],
            "rows": storage_rows,
        },
        "紧急购电量": {
            "headers": events.columns.tolist(),
            "rows": events.astype(object).where(events.notna(), None).values.tolist(),
        },
    }


def _result_interval_label(index: int) -> str:
    def format_minute(value: int) -> str:
        day, minute = divmod(value, 1440)
        hour, minute = divmod(minute, 60)
        return f"{hour}:{minute:02d}" + ("+1" if day else "")

    start = 10 * index
    return f"{format_minute(start)}-{format_minute(start + 10)}"


def _clear_data_rows(sheet) -> None:
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)


def _fill_purchase_sheet(sheet, rows, label: str) -> None:
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    try:
        total_column = headers.index("全天购电量") + 1
        cost_column = headers.index("全天购电费") + 1
    except ValueError as exc:
        raise ValueError(f"{label}缺少合计列") from exc

    interval_columns = list(range(2, total_column))
    if len(interval_columns) != T:
        raise ValueError(f"{label}应有144个时段列")
    canonical = [
        str(headers[column - 1]).replace("7:0-", "7:00-") for column in interval_columns
    ]
    if canonical[-1] == "0:00-0:10+1":
        canonical[-1] = "0:00+1-0:10+1"
    if canonical != [_result_interval_label(index) for index in range(1, T + 1)]:
        raise ValueError(f"{label}时段表头与模板约定不符")
    if sheet.max_row != len(rows) + 1:
        raise ValueError(f"{label}日期行数与结果不一致")

    template_dates = pd.DatetimeIndex(
        [pd.Timestamp(sheet.cell(row, 1).value) for row in range(2, sheet.max_row + 1)]
    ).normalize()
    result_dates = pd.DatetimeIndex([pd.Timestamp(row[0]) for row in rows]).normalize()
    if not template_dates.equals(result_dates):
        raise ValueError(f"{label}日期与模板不一致")

    for row_index, values in enumerate(rows, start=2):
        natural = np.asarray(values[1 : 1 + T], dtype=float)
        rotated = np.r_[natural[1:], natural[:1]]
        for column, value in zip(interval_columns, rotated):
            sheet.cell(row_index, column, float(max(0.0, value)))
        sheet.cell(row_index, total_column, float(values[-2]))
        sheet.cell(row_index, cost_column, float(values[-1]))


def write_result3(template_path: Path, output_path: Path, tables: dict) -> None:
    """从官方模板生成result3，只替换题目要求的数据区。"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path)
    required = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != required:
        raise ValueError(
            f"result3模板工作表应为{required}，实际为{workbook.sheetnames}"
        )

    for name in ("计划购电量", "调整购电量"):
        _fill_purchase_sheet(workbook[name], tables[name]["rows"], name)

    storage = workbook["充放电量"]
    storage_styles = [
        [copy(storage.cell(2 + block, column)._style) for column in range(1, 7)]
        for block in range(6)
    ]
    _clear_data_rows(storage)
    for index, values in enumerate(tables["充放电量"]["rows"]):
        target_row = index + 2
        block = index % 6
        for column in range(1, 7):
            storage.cell(target_row, column)._style = copy(
                storage_styles[block][column - 1]
            )
        for column, value in enumerate(values, start=1):
            if column == 1 and value is not None:
                value = pd.Timestamp(value).to_pydatetime()
            storage.cell(target_row, column, value)

    emergency = workbook["紧急购电量"]
    emergency_styles = [
        [copy(emergency.cell(row, column)._style) for column in range(1, 4)]
        for row in (2, 3, 4)
    ]
    emergency_date_format = emergency.cell(2, 1).number_format
    _clear_data_rows(emergency)
    events = pd.DataFrame(
        tables["紧急购电量"]["rows"],
        columns=tables["紧急购电量"]["headers"],
    )
    target_row = 2
    for date, group in events.groupby("日期", sort=True):
        items = list(group.itertuples(index=False, name=None))
        for item_index, (_, period, amount) in enumerate(items):
            if len(items) == 1 or item_index == len(items) - 1:
                style_index = 2
            elif item_index == 0:
                style_index = 0
            else:
                style_index = 1
            for column in range(1, 4):
                emergency.cell(target_row, column)._style = copy(
                    emergency_styles[style_index][column - 1]
                )
            emergency.cell(
                target_row,
                1,
                pd.Timestamp(date).to_pydatetime() if item_index == 0 else None,
            )
            if item_index == 0:
                emergency.cell(target_row, 1).number_format = emergency_date_format
            emergency.cell(target_row, 2, period)
            emergency.cell(target_row, 3, None if pd.isna(amount) else float(amount))
            target_row += 1

    workbook.save(output_path)


def write_revision_log(records, folder):
    """Monthly audit files: stream rows, then verify exact row counts before commit."""
    target = folder / "revision_logs"
    target.mkdir(parents=True, exist_ok=True)
    months = {}
    for row in records:
        months.setdefault(row["date"][:7], []).append(row)
    for month, rows in months.items():
        destination = target / f"q3_revisions_{month}.csv"
        temporary = destination.with_suffix(".tmp")
        with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        with temporary.open(encoding="utf-8-sig", newline="") as stream:
            count = sum(1 for _ in csv.DictReader(stream))
        if count != len(rows):
            raise RuntimeError(
                f"Incomplete revision export: {month}: {count}/{len(rows)}"
            )
        temporary.replace(destination)
    legacy = folder / "q3_revision_log.csv"
    if legacy.exists():
        legacy.unlink()


def save_case(folder, df, revisions, solves, cfg, price):
    folder.mkdir(parents=True, exist_ok=True)
    daily, events = daily_summary(df), emergency_events(df)
    df.to_csv(folder / "q3_timeseries.csv", index=False, encoding="utf-8-sig")
    daily.to_csv(folder / "q3_daily_summary.csv", index=False, encoding="utf-8-sig")
    events.to_csv(folder / "q3_emergency_events.csv", index=False, encoding="utf-8-sig")
    if len(revisions) != sum(T - int(row["issue_hour"]) * 6 for row in solves):
        raise RuntimeError("Missing forecast-revision records")
    write_revision_log(revisions, folder)
    pd.DataFrame(solves).to_csv(
        folder / "q3_solver_log.csv", index=False, encoding="utf-8-sig"
    )
    checks = audit(df, cfg)
    (folder / "q3_audit.json").write_text(
        json.dumps(checks, indent=2), encoding="utf-8"
    )
    tables = table_payload(df, daily, events)
    for name, obj in tables.items():
        pd.DataFrame(obj["rows"], columns=obj["headers"]).to_csv(
            folder / f"{name}.csv", index=False, encoding="utf-8-sig"
        )
    table1 = df[df.date.isin(SELECTED) & df.t.isin([60, 72, 84, 96, 108, 120])]
    table1[["date", "interval", "q0_kwh", "active_kwh", "emergency_kwh"]].to_csv(
        folder / "selected_table1.csv", index=False, encoding="utf-8-sig"
    )
    daily[daily.date.isin(SELECTED)].to_csv(
        folder / "selected_daily.csv", index=False, encoding="utf-8-sig"
    )
    storage = pd.DataFrame(
        tables["充放电量"]["rows"], columns=tables["充放电量"]["headers"]
    )
    storage["日期"] = storage["日期"].ffill()
    storage[storage["日期"].isin(SELECTED)].to_csv(
        folder / "selected_table2.csv", index=False, encoding="utf-8-sig"
    )
    events[events["日期"].isin(SELECTED)].to_csv(
        folder / "selected_table3.csv", index=False, encoding="utf-8-sig"
    )
    payload = {
        "tables": tables,
        "price": price.tolist(),
        "daily": daily.to_dict("records"),
        "config": asdict(cfg),
        "notes": [
            "Source: C题.pdf Q3/附录1-2; 附件1-3.xlsx. All energies kWh; costs yuan.",
            "Headers use natural day 00:00-24:00. Template shifted/ambiguous last interval corrected.",
            "计划购电量 fee = midnight plan fee; 调整购电量 fee = all-in actual settled fee.",
            "调整购电量 cells hold final active scheduled delivery, excluding emergency purchase.",
            "January: common midnight-only warm-up from 6000 kWh. Evaluation begins Feb 1.",
            "Results are chronological heuristic backtests, not a multistage optimum.",
        ],
    }
    (folder / "q3_workbook_payload.json").write_text(
        json.dumps(payload, ensure_ascii=False, allow_nan=False), encoding="utf-8"
    )
    return daily, checks, tables


def run(args):
    cfg = Config(
        settlement=args.settlement,
        eta_c=args.eta,
        eta_d=args.eta,
        plan_quantile=args.plan_quantile,
        update_quantile=args.update_quantile,
        reserve_weight=args.reserve_weight,
        load_bias_weight=args.load_bias_weight,
    )
    data = read_inputs(args.data_dir, args.end)
    cache, errors = make_forecast_cache(data, cfg)
    root = args.output_dir
    root.mkdir(parents=True, exist_ok=True)
    # ALL candidate policies share Jan operation and therefore Feb 1 inventory.
    e = E_INIT
    warm = []
    for i, date in enumerate(data.dates):
        if date >= pd.Timestamp("2025-02-01"):
            break
        e, frame, _, _ = simulate_day(data, cache, errors, i, e, (0,), cfg)
        warm.append(frame)
    warm_df = pd.concat(warm, ignore_index=True)
    audit(warm_df, cfg)
    warm_df.to_csv(root / "q3_january_warmup.csv", index=False, encoding="utf-8-sig")
    start_soc = e
    cases = (
        {"all": (0, 6, 12, 18)}
        if args.ablation
        else {
            (
                "all"
                if args.updates == "0,6,12,18"
                else "updates_" + args.updates.replace(",", "_")
            ): tuple(int(s) for s in args.updates.split(","))
        }
    )
    if args.ablation:
        cases.update(
            {"only_0": (0,), "0_6": (0, 6), "0_12": (0, 12), "0_6_12": (0, 6, 12)}
        )
    for hours in cases.values():
        if (
            not hours
            or hours[0] != 0
            or tuple(sorted(set(hours))) != hours
            or not set(hours) <= set(ISSUES)
        ):
            raise ValueError(
                "--updates must be a sorted subset of 0,6,12,18 starting with 0"
            )
    primary_case = next(iter(cases))
    primary_tables = None
    summaries = []
    for case, hours in cases.items():
        frames, revisions, solves, e = [], [], [], start_soc
        for di, date in enumerate(data.dates):
            if date < pd.Timestamp("2025-02-01"):
                continue
            e, df, rev, logs = simulate_day(data, cache, errors, di, e, hours, cfg)
            frames.append(df)
            revisions.extend(rev)
            solves.extend(logs)
            if date.day == 1 or date == data.dates[-1]:
                print(f"{case}: {date.date()} SOC={e:.2f}", flush=True)
        df = pd.concat(frames, ignore_index=True)
        daily, checks, tables = save_case(
            root / case, df, revisions, solves, cfg, data.price
        )
        if case == primary_case:
            primary_tables = tables
        raw_cost = float(daily.total_cost.sum())
        # Explicit reporting sensitivity, not a claimed unique economic salvage value.
        salvage_rate = cfg.eta_d * float(data.price.min())
        summaries.append(
            {
                "case": case,
                "updates": ",".join(map(str, hours)),
                "days": len(daily),
                "total_cost_yuan": raw_cost,
                "plan_cost_yuan": float(daily.plan_cost.sum()),
                "up_cost_yuan": float(daily.up_cost.sum()),
                "down_penalty_yuan": float(daily.down_penalty.sum()),
                "refund_yuan": float(daily.refund.sum()),
                "emergency_cost_yuan": float(daily.emergency_cost.sum()),
                "emergency_kwh": float(daily.emergency_kwh.sum()),
                "emergency_days": int((daily.emergency_kwh > 1e-6).sum()),
                "spill_kwh": float(daily.spill_kwh.sum()),
                "start_soc_kwh": start_soc,
                "end_soc_kwh": e,
                "salvage_rate_yuan_per_soc_kwh": salvage_rate,
                "inventory_adjusted_cost_yuan": raw_cost
                - salvage_rate * (e - start_soc),
                "max_solver_gap": max(row["mip_gap"] for row in solves),
                "max_solver_absolute_gap": max(
                    row["mip_absolute_gap"] for row in solves
                ),
                "audit_passed": checks["passed"],
            }
        )
    summary = pd.DataFrame(summaries)
    if args.ablation:
        base = summary.loc[summary.case == "only_0"].iloc[0]
        summary["saving_vs_only0_yuan"] = base.total_cost_yuan - summary.total_cost_yuan
        summary["inventory_adjusted_saving_yuan"] = (
            base.inventory_adjusted_cost_yuan - summary.inventory_adjusted_cost_yuan
        )
    summary.to_csv(root / "q3_comparison.csv", index=False, encoding="utf-8-sig")
    if data.dates[-1] == pd.Timestamp("2025-12-31"):
        write_result3(args.template, args.result, primary_tables)
        print(f"result3: {args.result.resolve()}", flush=True)
    else:
        print("短期测试未覆盖全年，跳过正式result3.xlsx。", flush=True)
    meta = {
        "config": asdict(cfg),
        "cases": cases,
        "evaluation_start": "2025-02-01",
        "evaluation_end": str(data.dates[-1].date()),
        "initial_soc_jan1": E_INIT,
        "initial_soc_feb1": start_soc,
        "python": sys.version,
        "numpy": np.__version__,
        "pandas": pd.__version__,
        "scipy": scipy.__version__,
        "source_sha256": {
            f: hashlib.sha256((args.data_dir / f).read_bytes()).hexdigest()
            for f in ("附件1.xlsx", "附件2.xlsx", "附件3.xlsx")
        },
        "model_class": "causal quantile-plus-energy-reserve rolling MILP heuristic",
        "settlement_reference": "final delivery vs same-day midnight plan, once per interval",
        "parameter_selection": "defaults fixed before annual ablation; no annual tuning performed",
    }
    (root / "q3_run_metadata.json").write_text(
        json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(summary.to_string(index=False), flush=True)


def parse_args():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--data-dir", type=Path, default=ATTACHMENT_DIR)
    p.add_argument("--template", type=Path, default=TEMPLATE_DIR / "result3.xlsx")
    p.add_argument("--result", type=Path, default=RESULT_DIR / "result3.xlsx")
    p.add_argument("--output-dir", type=Path, default=AUXILIARY_DIR)
    p.add_argument("--end", default="2025-12-31", help="short test e.g. 2025-02-03")
    p.add_argument("--updates", default="0,6,12,18")
    p.add_argument(
        "--ablation", action="store_true", help="five forecast-release combinations"
    )
    p.add_argument("--settlement", choices=["no_refund", "refund"], default="no_refund")
    p.add_argument(
        "--eta",
        type=float,
        default=0.9,
        help="each-way efficiency; sqrt(.9) for round trip .9",
    )
    p.add_argument("--plan-quantile", type=float, default=0.8)
    p.add_argument("--update-quantile", type=float, default=0.7)
    p.add_argument("--reserve-weight", type=float, default=1.0)
    p.add_argument("--load-bias-weight", type=float, default=0.0)
    a = p.parse_args()
    if not (
        0 < a.eta <= 1
        and 0 < a.plan_quantile < 1
        and 0 < a.update_quantile < 1
        and a.reserve_weight >= 0
        and 0 <= a.load_bias_weight <= 1
    ):
        p.error("Invalid efficiency/quantile/reserve/load-bias parameter")
    return a


if __name__ == "__main__":
    arguments = parse_args()
    if not arguments.template.exists():
        raise FileNotFoundError(f"找不到result3模板：{arguments.template.resolve()}")
    run(arguments)
